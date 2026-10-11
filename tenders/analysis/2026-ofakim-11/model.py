import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule

OUT = sys.argv[1]
F = 'Arial'
BLUE = Font(name=F, color='0000FF')
BLACK = Font(name=F)
GREEN = Font(name=F, color='008000')
BOLD = Font(name=F, bold=True)
TITLE = Font(name=F, bold=True, size=14)
HDR = Font(name=F, bold=True, color='FFFFFF')
HFILL = PatternFill('solid', fgColor='1F4E79')
YEL = PatternFill('solid', fgColor='FFFF00')
GREY = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
ILS = '#,##0 ₪;(#,##0 ₪);-'
ILS2 = '#,##0.00 ₪;(#,##0.00 ₪);-'
NUM = '#,##0;(#,##0);-'
PTS = '0.00;(0.00);-'
PCT = '0.0%;(0.0%);-'
WRAP = Alignment(wrap_text=True, vertical='top', horizontal='right')

wb = Workbook()


def sheet(name, widths):
    ws = wb.create_sheet(name)
    ws.sheet_view.rightToLeft = True
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    return ws


def header(ws, row, labels, col0=1):
    for i, h in enumerate(labels):
        c = ws.cell(row=row, column=col0 + i, value=h)
        c.font, c.fill, c.border = HDR, HFILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[row].height = 45


def put(ws, ref, val, font=BLACK, fmt=None, fill=None, note=None):
    c = ws[ref]
    c.value = val
    c.font = font
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if note:
        c.comment = Comment(note, 'מודל')
    return c


# ------------------------------------------------------------------ מקרא
ws = wb.active
ws.title = 'מקרא'
ws.sheet_view.rightToLeft = True
ws.column_dimensions['A'].width = 110
lines = [
    ('מודל כלכלי למכרזים – מכרז 11/2026 עיריית אופקים (מרכזיית טלפוניה)', TITLE),
    ('', BLACK),
    ('מטרה: לקבוע הצעת מחיר, לבדוק כדאיות ולהשוות מול מתחרים. המודל בנוי לשימוש חוזר – למכרז הבא מעדכנים את הסעיפים, הטווחים, המשקלים וההנחות.', BLACK),
    ('', BLACK),
    ('מוסכמות צבע:', BOLD),
    ('טקסט כחול = נתון קלט שמזינים ידנית (מחירים, כמויות, הצעות מתחרים).', BLUE),
    ('רקע צהוב = הנחת מפתח שצריך לאמת מול נתוני HOT בפועל (עלויות, הסתברויות).', BLACK),
    ('טקסט שחור = נוסחה. אין לשנות.', BLACK),
    ('טקסט ירוק = קישור מגיליון אחר.', GREEN),
    ('', BLACK),
    ('סדר עבודה מומלץ:', BOLD),
    ('1. הנחות – לעדכן עלויות אמיתיות (ציוד, פלטפורמה, תמיכה, דקות). ערכי ברירת המחדל הם הערכות בלבד.', BLACK),
    ('2. סעיפי מחיר – להזין את ההצעה שלנו והערכת הצעות המתחרים. עמודת "₪ לנקודה" מראה איפה הכי זול לקנות ניקוד.', BLACK),
    ('3. תוספות – סעיפים שאינם בשקלול: אין להם השפעה על הניקוד, ולכן ברירת המחדל היא המחיר המרבי.', BLACK),
    ('4. ניקוד ותחרות – להזין הערכת ציוני איכות. רואים מי מוביל ובכמה.', BLACK),
    ('5. תזרים – רווח ל-36 חודשים, לכל התקופה, ערך נוכחי (NPV), ותרחיש סיום מוקדם.', BLACK),
    ('6. החלטת הגשה – כרטיס ניקוד Go / No-Go.', BLACK),
    ('', BLACK),
    ('הנוסחה לניקוד מחיר במכרז: ניקוד סעיף = משקל × (ההצעה הנמוכה בסעיף ÷ ההצעה הנבחנת); סכום הסעיפים × 0.7.', BLACK),
    ('מחיר מחוץ לטווח מינימום–מקסימום פוסל את ההצעה (נספח ה\', הנחיה 2). המודל מסמן זאת באדום.', BLACK),
]
for i, (t, f) in enumerate(lines, start=1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = f
    c.alignment = Alignment(wrap_text=True, horizontal='right')
ws['A7'].fill = YEL

# ------------------------------------------------------------------ הנחות
ws = sheet('הנחות', {'A': 52, 'B': 16, 'C': 70})
ws['A1'] = 'הנחות עלות ותפעול'
ws['A1'].font = TITLE
header(ws, 3, ['הנחה', 'ערך', 'הסבר / מקור'])
A = [
    # row, label, value, fmt, note, key
    (4, 'עלות ציוד ורישוי מרכזייה מקומית + Call Center מקומי (חד-פעמי)', 25000, ILS, 'הערכה. לעדכן לפי הצעת יצרן/מפיץ.', True),
    (5, 'עלות הקמה, התקנה, הגדרות, ממשקים והדרכה (חד-פעמי)', 15000, ILS, 'הערכה: כ-100 שעות עבודה. כולל ממשק CRM, AD ודוא"ל.', True),
    (6, 'עלות הכנת הצעה, הדגמה ומסמכים (חד-פעמי, גם אם מפסידים)', 8000, ILS, 'הערכה: רכישת מסמכים 500 ₪ + זמן צוות + הדגמה של 120 דקות.', True),
    (7, 'מספר שלוחות', 300, NUM, 'מפרט סעיף 2.3 ו-3.12.', False),
    (8, 'עלות פלטפורמת ענן לשלוחה לחודש (עלות שולית)', 4, ILS2, 'הערכה לעלות שולית פנימית. מחיר שוק לשלוחה הוא 20–80 ₪ (ראה מדדי שוק).', True),
    (9, 'עלות רישוי Call Center לחודש (30 עמדות, שני המוקדים)', 600, ILS, 'הערכה. מחיר שוק לעמדת מוקד: 69–100 ₪ לחודש.', True),
    (10, 'עלות תמיכה, תחזוקה ו-SLA לחודש (כוח אדם)', 1500, ILS, 'הערכה: מוקד, טכנאי, ניהול שינויים ללא תשלום (מפרט 6.4).', True),
    (11, 'דקות שיחה בחודש', 40000, NUM, 'מפרט סעיף 2.6.', False),
    (12, 'עלות ממוצעת לדקה יוצאת', 0.03, '0.000 ₪', 'הערכה משוקללת נייח/סלולר. דמי קישוריות לסלולר הם כ-6 אג\' לדקה (לא מאומת לשנת 2026). אם רוב הדקות On-Net – נמוך יותר.', True),
    (13, 'ערבויות, ביטוח (כולל הרחבת סייבר) ו-ISO – לשנה', 3500, ILS, 'הערכה: עמלת ערבות ביצוע 20,000 ₪ + תוספת פרמיה.', True),
    (14, 'קנסות SLA צפויים לשנה', 2000, ILS, 'הערכה. טבלת הקנסות במפרט עמ\' 18.', True),
    (15, 'עלות טלפון בסיסי (8.1)', 230, ILS, 'הערכה. מחיר קמעונאי בבריטניה ל-T33G: כ-£64 לפני מע"מ.', True),
    (16, 'עלות טלפון מתקדם (8.2)', 450, ILS, 'הערכה. מחיר קמעונאי בבריטניה ל-T46U: כ-£134 לפני מע"מ – ייתכן שגבוה מהמחיר המרבי 500 ₪.', True),
    (17, 'שנות אופציה צפויות (0–4)', 2, NUM, 'הנחה. העירייה רשאית להאריך עד 4 שנים.', True),
    (18, 'שיעור היוון שנתי', 0.08, PCT, 'הנחה: עלות הון.', True),
    (19, 'אופן תשלום סעיף 1: 1 = אבני דרך 20%/80%, 2 = 20% ואז 80% ב-36 תשלומים', 1, '0', 'הסכם 8.2.2: העירייה רשאית לבחור באפשרות 2.', True),
    (20, 'חודש סיום מוקדם (0 = ללא סיום מוקדם)', 0, '0', 'הסכם 7.3: סיום בהודעה של 30 יום. תרחיש: להזין למשל 12.', True),
    (21, 'חודש עלייה לאוויר', 2, '0', 'מסמך א\' 2.4: 30 ימים מההודעה על חתימה.', True),
    (22, 'מקדם ניקוד מחיר', 0.7, '0.00', 'מסמך א\' 10.2: הניקוד הסופי מוכפל ב-70%.', False),
]
for r, label, val, fmt, note, key in A:
    put(ws, f'A{r}', label).alignment = WRAP
    put(ws, f'B{r}', val, BLUE, fmt, YEL if key else None)
    put(ws, f'C{r}', note).alignment = WRAP
ws['A24'] = 'נגזרות'
ws['A24'].font = BOLD
put(ws, 'A25', 'עלות שוטפת לחודש, ללא דקות')
put(ws, 'B25', '=B7*B8+B9+B10+(B13+B14)/12', fmt=ILS)
put(ws, 'C25', 'פלטפורמה + רישוי מוקד + תמיכה + ערבויות/ביטוח/קנסות')
put(ws, 'A26', 'עלות דקות לחודש')
put(ws, 'B26', '=B11*B12', fmt=ILS)
put(ws, 'C26', 'דקות × עלות לדקה')
put(ws, 'A27', 'סיום התקשרות צפוי (חודש)')
put(ws, 'B27', '=IF(B20>0,B20,36+12*B17)', fmt='0')
put(ws, 'C27', 'אם אין סיום מוקדם: 36 חודשים + שנות אופציה')

# ------------------------------------------------------------------ סעיפי מחיר
ws = sheet('סעיפי מחיר', {'A': 5, 'B': 34, 'C': 9, 'D': 11, 'E': 11, 'F': 8, 'G': 10, 'H': 10,
                          'I': 12, 'J': 12, 'K': 12, 'L': 12, 'M': 12, 'N': 12, 'O': 18,
                          'P': 10, 'Q': 10, 'R': 10, 'S': 10, 'T': 14, 'U': 13, 'V': 13, 'W': 13, 'X': 13})
ws['A1'] = 'סעיפי מחיר בשקלול (נספח ה\')'
ws['A1'].font = TITLE
ws['B2'] = 'מקדם ניקוד:'
ws['C2'] = "='הנחות'!B22"
ws['C2'].font = GREEN
header(ws, 4, ['#', 'סעיף', 'יחידה', 'מחיר מינימום', 'מחיר מקסימום', 'משקל', 'כמות ב-36 ח\'', 'כמות בשנות אופציה',
               'ההצעה שלנו', 'העלות שלנו ליחידה', 'מתחרה 1', 'מתחרה 2', 'מתחרה 3', 'ההצעה הנמוכה (A)',
               'בדיקת טווח', 'ניקוד שלנו', 'ניקוד מתחרה 1', 'ניקוד מתחרה 2', 'ניקוד מתחרה 3',
               '₪ לנקודה אחת (על כל החוזה)', 'הכנסה 36 ח\'', 'הכנסה אופציה', 'עלות 36 ח\'', 'רווח 36 ח\''])
items = [
    (5, 1, 'חבילה: מרכזייה מקומית + ענן + 2 מוקדים, כולל 36 ח\' שירות', 'קומפלט', 100000, 122000, 65, 1, '0',
     "='הנחות'!B4+'הנחות'!B5+36*'הנחות'!B25"),
    (6, 2, 'טלפון בסיסי (8.1)', 'יחידה', 250, 420, 3, 30, '0', "='הנחות'!B15"),
    (7, 3, 'טלפון מתקדם (8.2)', 'יחידה', 420, 500, 2, 10, '0', "='הנחות'!B16"),
    (8, 4, 'שירות, תמיכה ואחריות משנה רביעית', 'שנתי', 20000, 23000, 15, 0, "='הנחות'!B17", "=12*'הנחות'!B25"),
    (9, 5, 'בנק 40,000 דקות', 'חודשי', 1200, 2200, 15, 36, "=12*'הנחות'!B17", "='הנחות'!B26"),
]
for r, n, name, unit, mn, mx, w, q36, qopt, cost in items:
    put(ws, f'A{r}', n)
    put(ws, f'B{r}', name).alignment = WRAP
    put(ws, f'C{r}', unit)
    put(ws, f'D{r}', mn, BLUE, ILS)
    put(ws, f'E{r}', mx, BLUE, ILS)
    put(ws, f'F{r}', w, BLUE, '0')
    put(ws, f'G{r}', q36, BLUE if r in (6, 7) else BLACK, NUM, YEL if r in (6, 7) else None,
        'כמות לא ידועה – הנחה. ראו שאלת הבהרה 10.' if r in (6, 7) else None)
    put(ws, f'H{r}', qopt if isinstance(qopt, int) else qopt, GREEN if str(qopt).startswith('=') else BLACK, NUM)
    put(ws, f'I{r}', mn, BLUE, ILS, YEL)
    put(ws, f'J{r}', cost, GREEN, ILS)
    for col in 'KLM':
        put(ws, f'{col}{r}', mn, BLUE, ILS)
    put(ws, f'N{r}', f'=MIN(I{r},K{r},L{r},M{r})', fmt=ILS)
    put(ws, f'O{r}', f'=IF(OR(I{r}<D{r},I{r}>E{r}),"פסול – מחוץ לטווח","תקין")')
    put(ws, f'P{r}', f'=IF(I{r}>0,F{r}*N{r}/I{r}*$C$2,0)', fmt=PTS)
    put(ws, f'Q{r}', f'=IF(K{r}>0,F{r}*N{r}/K{r}*$C$2,0)', fmt=PTS)
    put(ws, f'R{r}', f'=IF(L{r}>0,F{r}*N{r}/L{r}*$C$2,0)', fmt=PTS)
    put(ws, f'S{r}', f'=IF(M{r}>0,F{r}*N{r}/M{r}*$C$2,0)', fmt=PTS)
    put(ws, f'T{r}', f'=IF(F{r}*N{r}>0,(G{r}+H{r})*I{r}^2/(F{r}*$C$2*N{r}),0)', fmt=ILS,
        note='כמה שקלים (על פני כל החוזה) עולה להוסיף נקודת ניקוד אחת בסעיף הזה, בהינתן ההצעה הנוכחית. נמוך = זול לקנות ניקוד.')
    put(ws, f'U{r}', f'=I{r}*G{r}', fmt=ILS)
    put(ws, f'V{r}', f'=I{r}*H{r}', fmt=ILS)
    put(ws, f'W{r}', f'=J{r}*G{r}', fmt=ILS)
    put(ws, f'X{r}', f'=U{r}-W{r}', fmt=ILS)
r = 10
put(ws, f'B{r}', 'סה"כ', BOLD)
for col in 'PQRSUVWX':
    put(ws, f'{col}{r}', f'=SUM({col}5:{col}9)', BOLD, PTS if col in 'PQRS' else ILS)
put(ws, 'B12', 'פער ניקוד מחיר מרבי אפשרי (כל הסעיפים במינימום מול כל הסעיפים במקסימום)', BOLD).alignment = WRAP
put(ws, 'P12', '=SUMPRODUCT(F5:F9,1-D5:D9/E5:E9)*$C$2', BOLD, PTS)
put(ws, 'B13', 'מחיר שלוחה לחודש שמגולם בסעיף 1 (ההצעה שלנו)').alignment = WRAP
put(ws, 'P13', "=I5/36/'הנחות'!B7", fmt=ILS2)
put(ws, 'B14', 'מחיר לדקה שמגולם בסעיף 5 (ההצעה שלנו)').alignment = WRAP
put(ws, 'P14', "=I9/'הנחות'!B11", fmt='0.000 ₪')
ws.conditional_formatting.add('O5:O9', CellIsRule(operator='equal', formula=['"פסול – מחוץ לטווח"'],
                                                  font=Font(name=F, color='FFFFFF', bold=True),
                                                  fill=PatternFill('solid', fgColor='C00000')))
ws.freeze_panes = 'C5'

# ------------------------------------------------------------------ תוספות
ws = sheet('תוספות', {'A': 5, 'B': 44, 'C': 11, 'D': 11, 'E': 12, 'F': 12, 'G': 12, 'H': 14, 'I': 18})
ws['A1'] = 'תוספות שאינן בשקלול (נספח ה\', עמ\' 57)'
ws['A1'].font = TITLE
ws['B2'] = 'אין להן השפעה על הניקוד – לכן ההמלצה היא המחיר המרבי. סעיף שלא ימולא יחייב בממוצע בין המינימום למקסימום.'
ws['B2'].alignment = WRAP
ws.row_dimensions[2].height = 30
header(ws, 4, ['#', 'סעיף', 'מינימום', 'מקסימום', 'ההצעה שלנו', 'העלות שלנו ליחידה', 'כמות צפויה לחודש',
               'רווח חודשי צפוי', 'בדיקת טווח'])
opts = [
    (1, 'שלוחה בודדת ב-Call Center בענן', 0, 45, 15, 0),
    (2, 'טלפון מתקדם שני מסכים (8.3)', 200, 780, 600, 0.2),
    (3, 'מתאם FXS עם 2 יציאות', 50, 350, 120, 0.2),
    (4, 'מתאם FXS עם 4 יציאות', 100, 600, 200, 0.1),
    (5, 'ספק כוח לטלפון IP', 20, 40, 15, 0.5),
    (6, 'שעת טכנאי', 250, 330, 150, 2),
    (7, 'רישוי 10 שלוחות נוספות ל-Call Center בענן', 300, 450, 200, 0),
    (8, 'רישוי 10 שלוחות נוספות ל-Call Center מקומי', 100, 150, 100, 0),
    (9, 'חבילת 10,000 דקות נוספות לחודש', 200, 350, 300, 0),
    (10, 'שעת הקלטה באולפן', 250, 330, 150, 0.1),
]
for i, (n, name, mn, mx, cost, qty) in enumerate(opts):
    r = 5 + i
    put(ws, f'A{r}', n)
    put(ws, f'B{r}', name)
    put(ws, f'C{r}', mn, BLUE, ILS)
    put(ws, f'D{r}', mx, BLUE, ILS)
    put(ws, f'E{r}', f'=D{r}', fmt=ILS, fill=YEL,
        note='ברירת מחדל: המחיר המרבי (אין השפעה על הניקוד). אפשר לדרוס בערך ידני.')
    put(ws, f'F{r}', cost, BLUE, ILS, YEL, 'הערכה.')
    put(ws, f'G{r}', qty, BLUE, '0.0', YEL, 'הערכה.')
    put(ws, f'H{r}', f'=(E{r}-F{r})*G{r}', fmt=ILS)
    put(ws, f'I{r}', f'=IF(OR(E{r}<C{r},E{r}>D{r}),"פסול – מחוץ לטווח","תקין")')
put(ws, 'B15', 'סה"כ רווח חודשי מתוספות', BOLD)
put(ws, 'H15', '=SUM(H5:H14)', BOLD, ILS)
put(ws, 'B16', 'הערה: עלות הטלפון שני המסכים עשויה לעלות על המינימום (200 ₪) – לא להתקרב למינימום.').alignment = WRAP
ws.conditional_formatting.add('I5:I14', CellIsRule(operator='equal', formula=['"פסול – מחוץ לטווח"'],
                                                   font=Font(name=F, color='FFFFFF', bold=True),
                                                   fill=PatternFill('solid', fgColor='C00000')))

# ------------------------------------------------------------------ ניקוד ותחרות
ws = sheet('ניקוד ותחרות', {'A': 18, 'B': 14, 'C': 16, 'D': 14, 'E': 14, 'F': 10, 'G': 22, 'H': 16})
ws['A1'] = 'ניקוד כולל ותחרות'
ws['A1'].font = TITLE
header(ws, 3, ['מציע', 'ניקוד מחיר (מתוך 70)', 'איכות הפתרון (מתוך 20)', 'מצגת (מתוך 10)', 'סה"כ', 'דירוג',
               'סף איכות 70% (14 נק\')', 'פער מול המוביל האחר'])
bidders = [('HOT (אנחנו)', 'P', 16, 8), ('מתחרה 1', 'Q', 15, 7), ('מתחרה 2', 'R', 14, 7), ('מתחרה 3', 'S', 12, 6)]
for i, (name, col, q, p) in enumerate(bidders):
    r = 4 + i
    put(ws, f'A{r}', name, BOLD if i == 0 else BLACK)
    put(ws, f'B{r}', f"='סעיפי מחיר'!{col}10", GREEN, PTS)
    put(ws, f'C{r}', q, BLUE, '0.0', YEL, 'הערכה סובייקטיבית.')
    put(ws, f'D{r}', p, BLUE, '0.0', YEL, 'הערכה סובייקטיבית.')
    put(ws, f'E{r}', f'=B{r}+C{r}+D{r}', BOLD, PTS)
    put(ws, f'F{r}', f'=RANK(E{r},$E$4:$E$7)', fmt='0')
    put(ws, f'G{r}', f'=IF(C{r}>=14,"עובר","בסיכון פסילה")')
others = {4: 'E5:E7', 5: 'E4,E6:E7', 6: 'E4:E5,E7', 7: 'E4:E6'}
for r, rng in others.items():
    put(ws, f'H{r}', f'=E{r}-MAX({rng})', fmt=PTS)
put(ws, 'A10', 'שווי נקודת ניקוד אחת', BOLD)
put(ws, 'A11', 'בסעיף 1 (חבילה)')
put(ws, 'B11', "='סעיפי מחיר'!T5", GREEN, ILS)
put(ws, 'A12', 'בסעיף 5 (דקות, כל החוזה)')
put(ws, 'B12', "='סעיפי מחיר'!T9", GREEN, ILS)
put(ws, 'C11', 'כלומר: נקודת איכות אחת שווה בערך הנחה כזו במחיר. השקעה בהדגמה ובמסמך המענה זולה בהרבה.').alignment = WRAP
ws.merge_cells('C11:H12')
ws.conditional_formatting.add('G4:G7', CellIsRule(operator='equal', formula=['"בסיכון פסילה"'],
                                                  font=Font(name=F, color='C00000', bold=True)))

# ------------------------------------------------------------------ תזרים
ws = sheet('תזרים', {c: 13 for c in 'ABCDEFGHIJKLM'})
ws.column_dimensions['A'].width = 8
ws['A1'] = 'תזרים חודשי ורווחיות'
ws['A1'].font = TITLE
S = "'הנחות'!"
P = "'סעיפי מחיר'!"
summary = [
    ('הכנסות 36 חודשים', '=SUM(C8:G44)', ILS),
    ('עלויות 36 חודשים', '=SUM(H8:J44)', ILS),
    ('רווח 36 חודשים', '=SUM(K8:K44)', ILS),
    ('רווח כל התקופה', '=SUM(K8:K92)', ILS),
    ('ערך נוכחי נקי (NPV)', f'=K8+NPV({S}B18/12,K9:K92)', ILS),
    ('עלות חודשית מרבית לאיזון ב-36 ח\'', '=(SUM(C8:G44)-SUM(H8:H44)-SUM(J8:J44))/36', ILS),
]
for i, (lab, f, fmt) in enumerate(summary):
    col = 'B' if i < 3 else 'H'
    rr = 2 + (i % 3)
    lc = 'A' if i < 3 else 'F'
    put(ws, f'{lc}{rr}', lab, BOLD)
    ws.merge_cells(f'{lc}{rr}:{chr(ord(lc)+1)}{rr}')
    put(ws, f'{chr(ord(lc)+2)}{rr}', f, BOLD, fmt)
put(ws, 'J4', f'=IF(H4<{S}B25,"העלות השוטפת גבוהה מנקודת האיזון","תקין")')
header(ws, 7, ['חודש', 'פעיל', 'הכנסה: סעיף 1', 'הכנסה: טלפונים', 'הכנסה: דקות', 'הכנסה: שירות שנה 4+',
               'רווח: תוספות', 'עלות חד-פעמית', 'עלות שוטפת', 'עלות דקות', 'תזרים נטו', 'מצטבר'])
GL = f'{S}$B$21'
for m in range(0, 85):
    r = 8 + m
    put(ws, f'A{r}', m, fmt='0')
    put(ws, f'B{r}', f'=IF(A{r}<={S}$B$27,1,0)', fmt='0')
    put(ws, f'C{r}', f'=IF({S}$B$19=1,IF(A{r}=1,0.2*{P}$I$5,0)+IF(A{r}={GL},0.8*{P}$I$5,0)*B{r},'
                     f'IF(A{r}=1,0.2*{P}$I$5,0)+IF(AND(A{r}>={GL},A{r}<{GL}+36),0.8*{P}$I$5/36,0)*B{r})', fmt=ILS)
    put(ws, f'D{r}', f'=IF(A{r}=1,{P}$I$6*{P}$G$6+{P}$I$7*{P}$G$7,0)', fmt=ILS)
    put(ws, f'E{r}', f'=IF(A{r}>={GL},{P}$I$9,0)*B{r}', fmt=ILS)
    put(ws, f'F{r}', f'=IF(A{r}>36,{P}$I$8/12,0)*B{r}', fmt=ILS)
    put(ws, f'G{r}', f"=IF(A{r}>={GL},'תוספות'!$H$15,0)*B{r}", fmt=ILS)
    put(ws, f'H{r}', f'=IF(A{r}=0,{S}$B$6,0)+IF(A{r}=1,{S}$B$4+{S}$B$5+{P}$J$6*{P}$G$6+{P}$J$7*{P}$G$7,0)', fmt=ILS)
    put(ws, f'I{r}', f'=IF(A{r}>=1,{S}$B$25,0)*B{r}', fmt=ILS)
    put(ws, f'J{r}', f'=IF(A{r}>={GL},{S}$B$26,0)*B{r}', fmt=ILS)
    put(ws, f'K{r}', f'=C{r}+D{r}+E{r}+F{r}+G{r}-H{r}-I{r}-J{r}', fmt=ILS)
    put(ws, f'L{r}', f'=K{r}' if m == 0 else f'=L{r-1}+K{r}', fmt=ILS)
put(ws, 'F5', 'חשיפת מזומנים מרבית', BOLD)
ws.merge_cells('F5:G5')
put(ws, 'H5', '=MIN(L8:L92)', BOLD, ILS)
put(ws, 'A5', 'הערה: באפשרות תשלום 2 ובסיום מוקדם, התשלומים שטרם שולמו אובדים (לפי נוסח ההסכם הנוכחי – ראו שאלת הבהרה 4).')
ws.merge_cells('A5:E6')
ws['A5'].alignment = WRAP
ws.freeze_panes = 'B8'

# ------------------------------------------------------------------ מדדי שוק
ws = sheet('מדדי שוק', {'A': 38, 'B': 22, 'C': 46, 'D': 60})
ws['A1'] = 'מדדי שוק (מחקר רשת, אוקטובר 2026)'
ws['A1'].font = TITLE
header(ws, 3, ['מדד', 'ערך', 'הערה / אמינות', 'מקור'])
bm = [
    ('מרכזייה בענן – מחיר לשלוחה לחודש (שוק SMB)', '20–80 ₪ לפני מע"מ', 'מדריך השוואה שיווקי. לפי אותו מקור רוב המרכזיות עד 25 ₪.', 'https://www.xn----9hcbbmdejomb6a2c8b8a.org.il/cloud-pbx/'),
    ('Voipe – משתמש בסיסי / משתמש Call Center', '39 ₪ / 69 ₪ לחודש', 'מחירון ספק, חיוב שנתי, מינימום 5 משתמשים.', 'https://voipe.co.il/pbx'),
    ('בזק עסקים IP Centrex', '26.90–37 ₪ לשלוחה', 'לפי מדריך השוואה; לא אומת מול מחירון בזק.', 'https://www.xn----9hcbbmdejomb6a2c8b8a.org.il/cloud-pbx/'),
    ('עמדת Call Center מתקדמת (CRM/AI)', 'עד כ-100–180 ₪ לחודש', 'מדריכי השוואה; טווח רחב.', 'https://www.teleco.co.il/%D7%9E%D7%A8%D7%9B%D7%96%D7%99%D7%95%D7%AA-%D7%91%D7%A2%D7%A0%D7%9F/'),
    ('Yealink T33G (בסיסי)', 'כ-£64 לפני מע"מ (בריטניה)', 'מחיר קמעונאי בחו"ל; מחיר מפיץ בישראל לא נמצא.', 'https://www.netxl.com/desktop-voip-phones/yealink-t33g-sip-entry-level-desk-phone-with-no-psu-included'),
    ('Yealink T46U (מתקדם)', 'כ-£134 לפני מע"מ (בריטניה)', 'מחיר קמעונאי בחו"ל. ייתכן שהמחיר המרבי במכרז (500 ₪) נמוך מהעלות.', 'https://www.provu.co.uk/yealink-t46u.html'),
    ('דמי קישוריות גומלין לסלולר', 'כ-6 אג\' לדקה', 'מתווה הביטול נעצר לפי דיווחים; התעריף העדכני לא אומת.', 'https://www.calcalist.co.il/local_news/article/hy7p11dsp2'),
    ('הצעה גרעונית – עמדת הפסיקה', 'לא נפסלת אוטומטית', 'עע"מ 3499/08 רון עבודות נ\' ועדת המכרזים עפולה; הצעה תכסיסנית – נפסלת.', 'https://www.duns100.co.il/articles/לא_כל_מספר_זוכה__על_הצעה_הפסדית_והצעה_תכסיסנית_במכרז_הציבורי'),
    ('מחיר שלוחה לחודש שמגולם בסעיף 1 באופקים', "='סעיפי מחיר'!P13", 'חישוב: מחיר החבילה ÷ 36 ÷ 300. כולל גם חומרה מקומית ושני מוקדים.', 'חישוב פנימי'),
    ('תוצאות מכרזי טלפוניה עירוניים קודמים', 'לא נמצאו מחירים פומביים', 'פרוטוקולי ועדות (למשל רמלה 15/16) אינם מפרסמים סכומים. מומלץ לבקש עיון במסמכי הזוכה.', 'https://www.ramle.org.il/uploads/n/1614840111.5536.pdf'),
]
for i, (a, b, c, d) in enumerate(bm):
    r = 4 + i
    put(ws, f'A{r}', a).alignment = WRAP
    cell = put(ws, f'B{r}', b, GREEN if str(b).startswith('=') else BLACK, ILS2 if str(b).startswith('=') else None)
    put(ws, f'C{r}', c).alignment = WRAP
    put(ws, f'D{r}', d).alignment = WRAP
    ws.row_dimensions[r].height = 42

# ------------------------------------------------------------------ החלטת הגשה
ws = sheet('החלטת הגשה', {'A': 44, 'B': 10, 'C': 12, 'D': 14, 'E': 60})
ws['A1'] = 'כרטיס החלטה: להגיש / לא להגיש'
ws['A1'].font = TITLE
header(ws, 3, ['קריטריון', 'משקל', 'ציון (1–5)', 'ציון משוקלל', 'הערה'])
crit = [
    ('עמידה בכל תנאי הסף (כולל 3 לקוחות × 50 שלוחות אנלוגיות)', 0.20, 3, 'אם לא עומדים – לא מגישים, בלי קשר לשאר.'),
    ('רווחיות צפויה (גיליון תזרים)', 0.20, 2, 'בהנחות ברירת המחדל הרווח שלילי – לעדכן עלויות.'),
    ('ערך אסטרטגי ואפסייל (קישוריות, סלולר, בתי ספר)', 0.15, 4, 'שלוחות SIM/eSIM ובתי ספר – הזדמנות ל-HOT.'),
    ('סיכון חוזי (קנסות, סיום מוקדם, אחריות ללא תקרה)', 0.15, 2, 'תלוי בתשובות להבהרות 2–4.'),
    ('יכולת הדגמה חיה ואיכות פתרון (30 נק\')', 0.15, 4, '30 נקודות איכות מול פער מחיר מרבי של כ-15 נקודות.'),
    ('עוצמת התחרות והספק המכהן', 0.10, 3, 'הספק הקיים (SAR) מכיר את הלקוח.'),
    ('זמינות צוות ומשאבים בלוח הזמנים', 0.05, 3, 'הגשה 03/11, הקמה תוך 30 יום.'),
]
for i, (name, w, s, note) in enumerate(crit):
    r = 4 + i
    put(ws, f'A{r}', name).alignment = WRAP
    put(ws, f'B{r}', w, BLUE, PCT)
    put(ws, f'C{r}', s, BLUE, '0', YEL)
    put(ws, f'D{r}', f'=B{r}*C{r}', fmt='0.00')
    put(ws, f'E{r}', note).alignment = WRAP
    ws.row_dimensions[r].height = 30
put(ws, 'A12', 'סה"כ (מתוך 5)', BOLD)
put(ws, 'B12', '=SUM(B4:B10)', BOLD, PCT)
put(ws, 'D12', '=SUM(D4:D10)', BOLD, '0.00')
put(ws, 'E12', '=IF(C4<3,"לא להגיש – תנאי סף",IF(D12>=3.5,"להגיש",IF(D12>=2.75,"להגיש בתנאים – לפי תשובות ההבהרה","לא להגיש")))', BOLD)
put(ws, 'A13', 'רווח צפוי ל-36 חודשים (מקושר)')
put(ws, 'D13', "='תזרים'!C4", GREEN, ILS)

wb.save(OUT)
print('saved', OUT)
