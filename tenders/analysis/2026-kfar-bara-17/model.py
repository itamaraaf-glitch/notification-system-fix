"""מודל כלכלי – מכרז מסגרת 17/2026, מועצה מקומית כפר ברא (שירותי IT מנוהלים).

הרצה:  python3 model.py <out.xlsx>   ואז חישוב מחדש (Excel / LibreOffice).
כל ההנחות בצהוב הן הערכות – לעדכן מול עלויות HOT בפועל.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule

OUT = sys.argv[1]
F = 'Arial'
BLUE, BLACK, GREEN = Font(name=F, color='0000FF'), Font(name=F), Font(name=F, color='008000')
BOLD, TITLE = Font(name=F, bold=True), Font(name=F, bold=True, size=14)
HDR, HFILL = Font(name=F, bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F4E79')
YEL = PatternFill('solid', fgColor='FFFF00')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
ILS, ILS2 = '#,##0 ₪;(#,##0 ₪);-', '#,##0.00 ₪;(#,##0.00 ₪);-'
NUM, PTS, PCT = '#,##0;(#,##0);-', '0.00;(0.00);-', '0.0%;(0.0%);-'
WRAP = Alignment(wrap_text=True, vertical='top', horizontal='right')
RED_RULE = dict(font=Font(name=F, color='FFFFFF', bold=True), fill=PatternFill('solid', fgColor='C00000'))

wb = Workbook()


def sheet(name, widths):
    ws = wb.create_sheet(name)
    ws.sheet_view.rightToLeft = True
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    return ws


def header(ws, row, labels, col0=1):
    for i, h in enumerate(labels):
        c = ws.cell(row=row, column=col0 + i, value=h)
        c.font, c.fill, c.border = HDR, HFILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[row].height = 48


def put(ws, ref, val, font=BLACK, fmt=None, fill=None, note=None):
    c = ws[ref]
    c.value, c.font, c.border = val, font, BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if note:
        c.comment = Comment(note, 'מודל')
    return c


H = "'הנחות'!"
B = "'רכיבי בסיס'!"
O = "'אופציונליים'!"

# ------------------------------------------------------------------ מקרא
ws = wb.active
ws.title = 'מקרא'
ws.sheet_view.rightToLeft = True
ws.column_dimensions['A'].width = 115
for i, (t, f) in enumerate([
    ('מודל כלכלי – מכרז מסגרת 17/2026, מועצה מקומית כפר ברא (שירותי IT מנוהלים)', TITLE),
    ('', BLACK),
    ('מבנה ההצעה במכרז: (1) הנחה אחידה על מחירים חודשיים מרביים של 14 רכיבי בסיס; (2) הנחה על רכיבים אופציונליים; (3) דמי הקמה חד-פעמיים בש"ח, בלי תקרה; (4) הנחה על שעת תמיכה (250 ₪) ועל מחירי מדף.', BLACK),
    ('ניקוד: מחיר 60 = 80% דמי הקמה + 20% מחיר שירות שנתי (נוסחת הזול ÷ הנבחן); איכות 40. העדפת ספק מקומי בפער של עד 5%.', BLACK),
    ('', BLACK),
    ('מוסכמות צבע:', BOLD),
    ('טקסט כחול = קלט שמזינים ידנית.', BLUE),
    ('רקע צהוב = הנחה שצריך לאמת מול נתוני HOT (עלויות, כמויות, הצעות מתחרים).', BLACK),
    ('טקסט שחור = נוסחה. טקסט ירוק = קישור מגיליון אחר.', BLACK),
    ('', BLACK),
    ('סדר עבודה:', BOLD),
    ('1. הנחות – אינפלציה, שער דולר, פיגור תשלום, תקופה, עלות הקמה.', BLACK),
    ('2. רכיבי בסיס – כמויות צפויות ועלות חודשית לכל רכיב. עמודת "מרווח" מראה רכיבים שמפסידים כבר במחיר המרבי.', BLACK),
    ('3. אופציונליים – כמות צפויה (ברירת מחדל 0) ועלות.', BLACK),
    ('4. ניקוד ותחרות – דמי הקמה ושיעור הנחה שלנו ושל המתחרים, ציוני איכות, ספק מקומי. כולל טבלת רגישות לדמי ההקמה.', BLACK),
    ('5. תזרים – 60 חודשים: רווח, NPV, חשיפה מרבית, שחיקת מרווח מאינפלציה (המחירים נומינליים ל-5 שנים).', BLACK),
], start=1):
    c = ws.cell(row=i, column=1, value=t)
    c.font, c.alignment = f, Alignment(wrap_text=True, horizontal='right')
ws['A8'].fill = YEL

# ------------------------------------------------------------------ הנחות
ws = sheet('הנחות', {'A': 56, 'B': 14, 'C': 72})
ws['A1'] = 'הנחות'
ws['A1'].font = TITLE
header(ws, 3, ['הנחה', 'ערך', 'הסבר / מקור'])
rows = [
    (4, 'שיעור הנחה על רכיבי הבסיס (ההצעה שלנו)', 0.0, PCT, 'נספח מ\' 3 פרק 1. 0% = המחיר המרבי.', True),
    (5, 'שיעור הנחה על רכיבים אופציונליים', 0.0, PCT, 'נספח מ\' 3 פרק 2. לא נכלל בשקלול (לבדוק בתשובות ההבהרה).', True),
    (6, 'דמי הקמה – שירותי בסיס (ההצעה שלנו, ₪)', 15000, ILS, 'נספח מ\' 3 "דמי הקמה" שורה 1. 80% מציון המחיר! ראו גיליון ניקוד.', True),
    (7, 'עלות ההקמה בפועל (מיגרציה, התקנות, הדרכות)', 30000, ILS, 'הערכה: כ-150 שעות + ציוד הקמה.', True),
    (8, 'עלות הכנת הצעה, ראיון ומסמכים', 6000, ILS, 'רכישת מסמכים 1,500 ₪ + זמן צוות.', True),
    (9, 'חודש עלייה לאוויר (מתחיל חיוב חודשי)', 2, '0', 'הנחה.', True),
    (10, 'שנות אופציה צפויות (0–2)', 1, '0', 'סעיף 13: 3 שנים + 2 × שנה.', True),
    (11, 'חודש סיום מוקדם (0 = ללא)', 0, '0', 'סעיף 13 (21.2): ביטול בהודעה של 30 יום.', True),
    (12, 'פיגור תשלום בחודשים', 2, '0', 'הסכם 19.ד–ה: עד 15 יום לאישור + שוטף 45.', True),
    (13, 'אינפלציה שנתית (על עלויות בש"ח)', 0.03, PCT, 'המחירים נומינליים ל-5 שנים (נספח מ\' 3) – העלויות עולות, ההכנסה לא.', True),
    (14, 'שינוי שנתי בעלויות הצמודות לדולר', 0.0, PCT, 'מחירוני Microsoft/ESET/Fortinet ושער הדולר. Microsoft העלתה מחירים ב-07/2026.', True),
    (15, 'שער דולר (₪)', 3.055, '0.000', 'בנק ישראל, 09/10/2026 (fx-rate.json בריפו).', False),
    (16, 'שיעור היוון שנתי', 0.08, PCT, 'עלות הון.', True),
    (17, 'שעות תמיכה לפי קריאה בחודש', 4, '0', 'הנחה. מעבר לביקור השבועי הכלול ברכיב 1.', True),
    (18, 'הנחה על שעת תמיכה (מ-250 ₪)', 0.0, PCT, 'נספח מ\' 3 פרק "תמיכה טכנית".', True),
    (19, 'עלות שעת טכנאי', 150, ILS, 'הערכה, כולל נסיעה.', True),
    (20, 'רכש מוצרים לפי מחירי מדף – לחודש', 3000, ILS, 'הערכה.', True),
    (21, 'הנחה ממחירי המדף שנציע', 0.05, PCT, 'נספח מ\' 3 סעיף 7.', True),
    (22, 'מרווח רכש מול מחיר מדף (לפני הנחה)', 0.15, PCT, 'הערכה: מחיר הספק נמוך ב-15% ממחיר המדף.', True),
]
for r, lab, val, fmt, note, key in rows:
    put(ws, f'A{r}', lab).alignment = WRAP
    put(ws, f'B{r}', val, BLUE, fmt, YEL if key else None)
    put(ws, f'C{r}', note).alignment = WRAP
put(ws, 'A24', 'נגזרות', BOLD)
put(ws, 'A25', 'חודש סיום ההתקשרות')
put(ws, 'B25', '=IF(B11>0,B11,36+12*B10)', fmt='0')
put(ws, 'A26', 'הכנסה חודשית מתמיכה ומרכש')
put(ws, 'B26', '=B17*250*(1-B18)+B20*(1-B21)', fmt=ILS)
put(ws, 'A27', 'עלות חודשית לתמיכה ורכש')
put(ws, 'B27', '=B17*B19+B20*(1-B22)', fmt=ILS)

# ------------------------------------------------------------------ רכיבי בסיס
ws = sheet('רכיבי בסיס', {'A': 5, 'B': 30, 'C': 11, 'D': 9, 'E': 10, 'F': 11, 'G': 10, 'H': 12, 'I': 13, 'J': 13, 'K': 13, 'L': 11, 'M': 42})
ws['A1'] = 'רכיבי שירות בסיס חודשי (נספח מ\' 3 פרק 1)'
ws['A1'].font = TITLE
header(ws, 3, ['#', 'רכיב', 'מחיר חודשי מרבי ליחידה', 'כמות בטופס', 'כמות צפויה', 'עלות שלנו ליחידה לחודש', 'חלק צמוד דולר', 'מחיר אחרי הנחה',
               'הכנסה חודשית', 'עלות חודשית', 'רווח חודשי', 'מרווח', 'הערה / הנחת עלות'])
base = [
    (1, 'שירות IT מנוהל מקצה לקצה (כולל ביקור שבועי 8 ש\')', 4500, 1, 1, 4200, 0, 'ביקור שבועי: 4.33 × 8 ש\' × ~120 ₪ = 4,160 ₪, ועוד מוקד 24/7. כנראה הפסדי.'),
    (2, 'אחסון שרת DC (2 לשרידות)', 675, 2, 2, 300, 0, 'מכונה וירטואלית + רישוי Windows Server.'),
    (3, 'אחסון שרת FS 5TB (2 ב-DFS, NVMe)', 200, 2, 2, 450, 0, '5TB NVMe ב-200 ₪ לחודש – כנראה מתחת לעלות.'),
    (4, 'אחסון שרת TS (עד 35 משתמשים)', 800, 1, 1, 1000, 0.7, 'כולל רישוי RDS ל-35 משתמשים (SPLA, צמוד דולר) – כנראה מתחת לעלות.'),
    (5, 'אבטחת מידע לתחנות (ESET)', 25, 1, 35, 12, 1, 'לפי תחנה. כמות צפויה: הנחה.'),
    (6, 'אחסון שרת הדפסות', 540, 1, 1, 200, 0, ''),
    (7, 'ניהול ציוד תקשורת (מתגים, Firewall, אלחוטי)', 900, 1, 1, 400, 0, ''),
    (8, 'מרכזיית IP – 35 שלוחות, 10,000 דקות, הקלטה, אפליקציה', 2000, 1, 1, 700, 0, 'פלטפורמת HOT + דקות.'),
    (9, 'גיבוי יומי לאתר מרוחק (כלל השרתים)', 1350, 1, 1, 500, 0, ''),
    (10, 'Microsoft 365 Business Standard', 45, 1, 35, "=14*'הנחות'!B15", 1, 'מחירון $14 למשתמש מ-07/2026 × שער. מרווח של כ-5% בלבד.'),
    (11, 'קו מטרו IPVPN 300/150, כולל VDOM ו-2 חוות', 1600, 1, 1, 700, 0, 'עלות On-Net של HOT – הערכה.'),
    (12, 'שכירות מכשיר טלפון IP מתקדם', 30, 1, 35, 15, 0.5, 'מכשיר ~450 ₪ ÷ 36 + תחזוקה.'),
    (13, 'שכירות מתג מנוהל 48 PoE', 270, 1, 2, 260, 0.5, 'מתג ~8,000 ₪ ÷ 36 + תמיכה. מרווח דק.'),
    (14, 'שכירות נקודת גישה מנוהלת', 90, 1, 6, 65, 0.5, 'AP ~1,800 ₪ ÷ 36 + רישוי.'),
]
for i, (n, name, mx, fq, eq, cost, usd, note) in enumerate(base):
    r = 4 + i
    put(ws, f'A{r}', n)
    put(ws, f'B{r}', name).alignment = WRAP
    put(ws, f'C{r}', mx, BLUE, ILS)
    put(ws, f'D{r}', fq, BLUE, NUM)
    put(ws, f'E{r}', eq, BLUE, NUM, YEL if eq != fq else None, 'הנחה – כמות לא ידועה (שאלות הבהרה 6, 40).' if eq != fq else None)
    put(ws, f'F{r}', cost, GREEN if str(cost).startswith('=') else BLUE, ILS2, YEL)
    put(ws, f'G{r}', usd, BLUE, PCT)
    put(ws, f'H{r}', f"=C{r}*(1-{H}$B$4)", fmt=ILS2)
    put(ws, f'I{r}', f'=H{r}*E{r}', fmt=ILS)
    put(ws, f'J{r}', f'=F{r}*E{r}', fmt=ILS)
    put(ws, f'K{r}', f'=I{r}-J{r}', fmt=ILS)
    put(ws, f'L{r}', f'=IF(I{r}>0,K{r}/I{r},0)', fmt=PCT)
    put(ws, f'M{r}', note).alignment = WRAP
    ws.row_dimensions[r].height = 30
put(ws, 'B18', 'סה"כ חודשי', BOLD)
for col in 'IJK':
    put(ws, f'{col}18', f'=SUM({col}4:{col}17)', BOLD, ILS)
put(ws, 'L18', '=IF(I18>0,K18/I18,0)', BOLD, PCT)
put(ws, 'B19', 'מזה עלות צמודת דולר')
put(ws, 'J19', '=SUMPRODUCT(J4:J17,G4:G17)', fmt=ILS)
put(ws, 'B20', 'מזה עלות שקלית')
put(ws, 'J20', '=J18-J19', fmt=ILS)
put(ws, 'B22', 'מחיר שירות שנתי לשקלול (12 × מחיר אחרי הנחה × כמות בטופס)').alignment = WRAP
put(ws, 'I22', '=12*SUMPRODUCT(H4:H17,D4:D17)', BOLD, ILS)
put(ws, 'B23', 'מחיר שירות שנתי לפי מחירים מרביים (לפני הנחה)').alignment = WRAP
put(ws, 'I23', '=12*SUMPRODUCT(C4:C17,D4:D17)', fmt=ILS)
ws.conditional_formatting.add('L4:L17', CellIsRule(operator='lessThan', formula=['0'], **RED_RULE))

# ------------------------------------------------------------------ אופציונליים
ws = sheet('אופציונליים', {'A': 5, 'B': 34, 'C': 10, 'D': 13, 'E': 11, 'F': 12, 'G': 12, 'H': 12, 'I': 11, 'J': 34})
ws['A1'] = 'רכיבים אופציונליים (נספח מ\' 3 פרק 2) – ברירת מחדל: לא מוזמנים'
ws['A1'].font = TITLE
header(ws, 3, ['#', 'רכיב', 'יחידה', 'מחיר חודשי בטופס', 'כמות צפויה', 'עלות שלנו ליחידה', 'מחיר אחרי הנחה', 'רווח חודשי', 'מרווח', 'הערה'])
opt = [
    (1, 'Microsoft 365 Business Premium', 'משתמש', 85, "=22*'הנחות'!B15", 'מחירון $22.'),
    (2, 'Compliance / Purview for BP', 'משתמש', 38, 20, ''),
    (3, 'Exchange Online Archive', 'משתמש', 16, 10, ''),
    (4, 'Cloud-to-Cloud Backup – Mail', 'תיבה', 20, 8, ''),
    ('–', 'קו תקשורת גיבוי לאתר הראשי', 'חודשי', 540, 250, 'ללא מספר בטופס.'),
    (7, 'אחסון נוסף Microsoft 365', '1TB', 270, 150, ''),
    (9, 'שעת תמיכה באתר', 'שעה', 252, 150, 'סותר את 250 ₪ בפרק התמיכה.'),
    (11, 'אחסון שרת Collector ל-SOC/SIEM', 'שרת', 630, 300, ''),
    (14, 'SOC/SIEM מנוהל', 'חודשי', 6300, 4500, '24/7 או שעות פעילות – לא ברור.'),
    (15, 'FortiNAC Cloud או שווה ערך', 'מערכת', 1800, 1100, ''),
    (16, 'FortiAnalyzer Cloud או שווה ערך', 'מערכת', 630, 400, ''),
    (17, 'FortiManager Cloud או שווה ערך', 'מערכת', 720, 450, ''),
    (18, 'NTA / NPM', 'מערכת', 630, 350, ''),
    (19, 'שלוחה נוספת למרכזייה', 'שלוחה', 25, 8, ''),
    (20, 'Call Center', 'מערכת', 450, 200, ''),
    (21, 'מומחה אבטחת מידע (CISO)', 'חודשי', 2500, 1800, ''),
]
for i, (n, name, unit, price, cost, note) in enumerate(opt):
    r = 4 + i
    put(ws, f'A{r}', n)
    put(ws, f'B{r}', name).alignment = WRAP
    put(ws, f'C{r}', unit)
    put(ws, f'D{r}', price, BLUE, ILS)
    put(ws, f'E{r}', 0, BLUE, NUM, YEL)
    put(ws, f'F{r}', cost, GREEN if str(cost).startswith('=') else BLUE, ILS2, YEL)
    put(ws, f'G{r}', f"=D{r}*(1-{H}$B$5)", fmt=ILS2)
    put(ws, f'H{r}', f'=(G{r}-F{r})*E{r}', fmt=ILS)
    put(ws, f'I{r}', f'=IF(G{r}>0,(G{r}-F{r})/G{r},0)', fmt=PCT)
    put(ws, f'J{r}', note).alignment = WRAP
put(ws, 'B21', 'סה"כ', BOLD)
put(ws, 'G21', '=SUMPRODUCT(G4:G19,E4:E19)', BOLD, ILS)
put(ws, 'H21', '=SUM(H4:H19)', BOLD, ILS)
put(ws, 'B22', 'עלות חודשית')
put(ws, 'H22', '=G21-H21', fmt=ILS)
ws.conditional_formatting.add('I4:I19', CellIsRule(operator='lessThan', formula=['0'], **RED_RULE))

# ------------------------------------------------------------------ ניקוד ותחרות
ws = sheet('ניקוד ותחרות', {'A': 18, 'B': 13, 'C': 11, 'D': 14, 'E': 11, 'F': 11, 'G': 11, 'H': 10, 'I': 10, 'J': 10, 'K': 11, 'L': 10, 'M': 9, 'N': 12})
ws['A1'] = 'ניקוד ותחרות'
ws['A1'].font = TITLE
ws['A2'] = 'מחיר 60 = 0.6 × (80 × הקמה זולה ÷ הקמה נבחנת + 20 × שירות שנתי זול ÷ שירות שנתי נבחן). איכות 40.'
header(ws, 4, ['מציע', 'דמי הקמה (₪)', 'הנחה על הבסיס', 'מחיר שירות שנתי', 'ציון הקמה (מ-80)', 'ציון שירות (מ-20)', 'ציון מחיר (מ-60)',
               'ניסיון (מ-20)', 'ראיון (מ-10)', 'המלצות (מ-10)', 'סה"כ', 'ספק מקומי (1/0)', 'דירוג', 'זוכה צפוי'])
bidders = [('HOT (אנחנו)', f'={H}B6', f'={H}B4', 12, 8, 8, 0),
           ('מתחרה 1', 20000, 0.10, 14, 7, 8, 0),
           ('מתחרה 2', 10000, 0.15, 10, 6, 6, 0),
           ('מתחרה 3 (מקומי?)', 30000, 0.05, 6, 8, 6, 1)]
for i, (name, setup, disc, q1, q2, q3, local) in enumerate(bidders):
    r = 5 + i
    put(ws, f'A{r}', name, BOLD if i == 0 else BLACK)
    put(ws, f'B{r}', setup, GREEN if i == 0 else BLUE, ILS, None if i == 0 else YEL)
    put(ws, f'C{r}', disc, GREEN if i == 0 else BLUE, PCT, None if i == 0 else YEL)
    put(ws, f'D{r}', f"={B}$I$23*(1-C{r})", fmt=ILS)
    put(ws, f'E{r}', f'=IF(B{r}>0,80*MIN($B$5:$B$8)/B{r},0)', fmt=PTS)
    put(ws, f'F{r}', f'=IF(D{r}>0,20*MIN($D$5:$D$8)/D{r},0)', fmt=PTS)
    put(ws, f'G{r}', f'=0.6*(E{r}+F{r})', fmt=PTS)
    put(ws, f'H{r}', q1, BLUE, '0.0', YEL)
    put(ws, f'I{r}', q2, BLUE, '0.0', YEL)
    put(ws, f'J{r}', q3, BLUE, '0.0', YEL)
    put(ws, f'K{r}', f'=G{r}+H{r}+I{r}+J{r}', BOLD, PTS)
    put(ws, f'L{r}', local, BLUE, '0')
    put(ws, f'M{r}', f'=RANK(K{r},$K$5:$K$8)', fmt='0')
# local preference: a local bidder within 5% of the top score wins
put(ws, 'A10', 'הציון הגבוה ביותר', BOLD)
put(ws, 'K10', '=MAX(K5:K8)', BOLD, PTS)
put(ws, 'A11', 'ציון ספק מקומי מוביל')
put(ws, 'K11', '=SUMPRODUCT(MAX(K5:K8*L5:L8))', fmt=PTS)
put(ws, 'A12', 'ספק מקומי זוכה בהעדפה?')
put(ws, 'K12', '=IF(AND(K11>0,K11>=0.95*K10),"כן","לא")')
for i in range(4):
    r = 5 + i
    put(ws, f'N{r}', f'=IF($K$12="כן",IF(AND(L{r}=1,K{r}=$K$11),"זוכה (מקומי)",""),IF(M{r}=1,"זוכה",""))')
ws.conditional_formatting.add('N5:N8', CellIsRule(operator='notEqual', formula=['""'], font=Font(name=F, bold=True, color='006100'),
                                                  fill=PatternFill('solid', fgColor='C6EFCE')))

put(ws, 'A14', 'מחיר נקודה אחת (₪)', BOLD)
put(ws, 'A15', 'דרך דמי ההקמה')
put(ws, 'B15', '=B5/48', fmt=ILS,
    note='כשאנחנו הזולים: ציון הקמה = 48 × A/B. נקודה אחת = B/48 ₪, תשלום חד-פעמי.')
put(ws, 'A16', 'דרך הנחה על הבסיס (על כל החוזה)')
put(ws, 'B16', f"=D5/12*{H}B25/12", fmt=ILS,
    note='ציון שירות = 12 × A/B. נקודה אחת = B/12 ₪ לשנה, כפול מספר השנים.')
put(ws, 'C15', '=IF(B15>0,B16/B15,0)', BOLD, '0"x"')
put(ws, 'D15', 'פי כמה זול יותר לקנות ניקוד דרך דמי ההקמה').alignment = WRAP
ws.merge_cells('D15:G16')

put(ws, 'A18', 'רגישות: דמי ההקמה שלנו מול הציון שלנו', BOLD)
header(ws, 19, ['דמי הקמה שנבחן', 'ציון הקמה', 'ציון מחיר', 'ציון כולל', 'פער מול המוביל האחר', 'הכנסה מוותרת מול 15,000 ₪'])
for j, v in enumerate([1, 1000, 5000, 10000, 15000, 20000, 30000, 45000]):
    r = 20 + j
    put(ws, f'A{r}', v, BLUE, ILS)
    put(ws, f'B{r}', f'=80*MIN(A{r},MIN($B$6:$B$8))/A{r}', fmt=PTS)
    put(ws, f'C{r}', f'=0.6*(B{r}+$F$5)', fmt=PTS)
    put(ws, f'D{r}', f'=C{r}+$H$5+$I$5+$J$5', BOLD, PTS)
    put(ws, f'E{r}', f'=D{r}-MAX(0.6*(80*MIN(A{r},MIN($B$6:$B$8))/$B$6+$F$6)+$H$6+$I$6+$J$6,'
                     f'0.6*(80*MIN(A{r},MIN($B$6:$B$8))/$B$7+$F$7)+$H$7+$I$7+$J$7,'
                     f'0.6*(80*MIN(A{r},MIN($B$6:$B$8))/$B$8+$F$8)+$H$8+$I$8+$J$8)', fmt=PTS)
    put(ws, f'F{r}', f'=15000-A{r}', fmt=ILS)
put(ws, 'A29', 'שימו לב: בנוסחת "הזול ÷ הנבחן" הצעת הקמה סמלית מוחקת כמעט את כל 48 הנקודות של המתחרים. הוועדה רשאית לפסול הצעה "בלתי סבירה" (סעיף 7.9), ולכן זה גם סיכון וגם הזדמנות.').alignment = WRAP
ws.merge_cells('A29:N30')
ws.row_dimensions[29].height = 30

# ------------------------------------------------------------------ תזרים
ws = sheet('תזרים', {c: 13 for c in 'ABCDEFGHIJKLM'})
ws.column_dimensions['A'].width = 8
ws['A1'] = 'תזרים חודשי (60 חודשים)'
ws['A1'].font = TITLE
HDR_ROW, R0 = 8, 9
LAST = R0 + 60
rng = lambda c: f'{c}{R0}:{c}{LAST}'
end36 = R0 + 36
summary = [
    ('A2', 'רווח (צבירה) 36 חודשים', 'C2', f'=SUM(I{R0}:I{end36})'),
    ('A3', 'רווח (צבירה) כל התקופה', 'C3', f'=SUM({rng("I")})'),
    ('A4', 'ערך נוכחי נקי (NPV) של התזרים', 'C4', f"=L{R0}+NPV({H}B16/12,L{R0+1}:L{LAST})"),
    ('F2', 'חשיפת מזומנים מרבית', 'H2', f'=MIN({rng("M")})'),
    ('F3', 'מרווח חודשי בחודש העלייה לאוויר', 'H3', f"=INDEX({rng('C')},{H}B9+1)+INDEX({rng('D')},{H}B9+1)+INDEX({rng('E')},{H}B9+1)-INDEX({rng('G')},{H}B9+1)"),
    ('F4', 'מרווח חודשי בחודש האחרון', 'H4', f"=INDEX({rng('C')},{H}B25+1)+INDEX({rng('D')},{H}B25+1)+INDEX({rng('E')},{H}B25+1)-INDEX({rng('G')},{H}B25+1)"),
]
for lc, lab, vc, f in summary:
    put(ws, lc, lab, BOLD)
    ws.merge_cells(f'{lc}:{chr(ord(lc[0]) + 1)}{lc[1:]}')
    put(ws, vc, f, BOLD, ILS)
put(ws, 'A5', 'שחיקת מרווח: המחירים נומינליים ל-5 שנים, העלויות עולות לפי אינפלציה ומחירוני דולר.')
ws.merge_cells('A5:H5')
header(ws, HDR_ROW, ['חודש', 'פעיל', 'הכנסה: בסיס', 'הכנסה: אופציונליים', 'הכנסה: תמיכה ורכש', 'הכנסה: הקמה', 'עלות חודשית',
                     'עלות חד-פעמית', 'רווח (צבירה)', 'חיוב כולל', 'תקבול (אחרי פיגור)', 'תזרים נטו', 'מצטבר'])
for m in range(0, 61):
    r = R0 + m
    on = f'B{r}*(A{r}>={H}$B$9)'
    put(ws, f'A{r}', m, fmt='0')
    put(ws, f'B{r}', f'=IF(AND(A{r}>=1,A{r}<={H}$B$25),1,0)', fmt='0')
    put(ws, f'C{r}', f'={B}$I$18*{on}', fmt=ILS)
    put(ws, f'D{r}', f'={O}$G$21*{on}', fmt=ILS)
    put(ws, f'E{r}', f'={H}$B$26*{on}', fmt=ILS)
    put(ws, f'F{r}', f'=IF(A{r}={H}$B$9,{H}$B$6,0)', fmt=ILS)
    put(ws, f'G{r}', f'=({B}$J$20*(1+{H}$B$13)^((A{r}-1)/12)+{B}$J$19*(1+{H}$B$14)^((A{r}-1)/12)'
                     f'+{O}$H$22+{H}$B$27)*{on}', fmt=ILS)
    put(ws, f'H{r}', f'=IF(A{r}=0,{H}$B$8,0)+IF(A{r}=1,{H}$B$7,0)', fmt=ILS)
    put(ws, f'I{r}', f'=C{r}+D{r}+E{r}+F{r}-G{r}-H{r}', fmt=ILS)
    put(ws, f'J{r}', f'=C{r}+D{r}+E{r}+F{r}', fmt=ILS)
    put(ws, f'K{r}', f'=IF(A{r}-{H}$B$12>=0,INDEX($J${R0}:$J${LAST},A{r}-{H}$B$12+1),0)', fmt=ILS)
    put(ws, f'L{r}', f'=K{r}-G{r}-H{r}', fmt=ILS)
    put(ws, f'M{r}', f'=L{r}' if m == 0 else f'=M{r-1}+L{r}', fmt=ILS)
put(ws, f'A{LAST + 2}', 'הערה: חיובים שנצברו עד חודש הסיום ייגבו גם אחריו (פיגור תשלום) – לכן התזרים נמשך עד חודש 60 גם בסיום מוקדם.')
ws.freeze_panes = f'B{R0}'

# ------------------------------------------------------------------ מדדי שוק
ws = sheet('מדדי שוק', {'A': 40, 'B': 22, 'C': 50, 'D': 60})
ws['A1'] = 'מדדי שוק ונתונים חיצוניים'
ws['A1'].font = TITLE
header(ws, 3, ['מדד', 'ערך', 'הערה', 'מקור'])
for i, (a, b, c, d) in enumerate([
    ('Microsoft 365 Business Standard', '$14 למשתמש לחודש', 'מ-1.7.2026 (היה $12.50). מחיר מרבי בטופס: 45 ₪ ≈ $14.7.', 'https://lazyadmin.nl/office-365/microsoft-365-business-standard-vs-premium-2026-update/'),
    ('Microsoft 365 Business Premium', '$22 למשתמש לחודש', 'ללא שינוי ב-2026. מחיר בטופס: 85 ₪.', 'https://www.stmicro.net/blog/microsoft-365-price-increase-2026/'),
    ('שער דולר יציג', '3.055 ₪', '09/10/2026', 'בנק ישראל (fx-rate.json בריפו)'),
    ('מרכזייה בענן – שלוחה לחודש (שוק)', '20–80 ₪', 'בטופס: 2,000 ₪ ל-35 שלוחות + 10,000 דקות ≈ 57 ₪ לשלוחה.', 'https://www.xn----9hcbbmdejomb6a2c8b8a.org.il/cloud-pbx/'),
    ('הצעה הפסדית', 'לא נפסלת אוטומטית', 'עע"מ 3499/08; הצעה תכסיסנית או בלתי סבירה – נפסלת (ראו גם סעיף 7.9 במכרז).', 'https://www.duns100.co.il/articles/לא_כל_מספר_זוכה__על_הצעה_הפסדית_והצעה_תכסיסנית_במכרז_הציבורי'),
]):
    r = 4 + i
    for col, v in zip('ABCD', (a, b, c, d)):
        put(ws, f'{col}{r}', v).alignment = WRAP
    ws.row_dimensions[r].height = 36

# ------------------------------------------------------------------ החלטת הגשה
ws = sheet('החלטת הגשה', {'A': 46, 'B': 10, 'C': 11, 'D': 13, 'E': 62})
ws['A1'] = 'כרטיס החלטה: להגיש / לא להגיש'
ws['A1'].font = TITLE
header(ws, 3, ['קריטריון', 'משקל', 'ציון (1–5)', 'ציון משוקלל', 'הערה'])
for i, (name, w, s, note) in enumerate([
    ('תנאי סף: 5 שנות IT מנוהל ל-2 רשויות, ISO 9001 + 27001, תמיכה 24/7', 0.20, 3, 'אם לא – רק דרך שותף/קבלן משנה (שאלת הבהרה 11).'),
    ('רווחיות (גיליון תזרים)', 0.20, 3, 'רווח חיובי בהנחות ברירת המחדל, אבל רכיבים 3–4 מתחת לעלות ויש שחיקה מאינפלציה.'),
    ('ערך אסטרטגי (קווי מטרו, מרכזייה, אפסייל אבטחה)', 0.15, 4, 'רכיבים 8 ו-11 הם הליבה של HOT.'),
    ('סיכון חוזי (ללא הצמדה, ביטול ב-30 יום, אין כמויות)', 0.15, 2, 'תלוי בתשובות להבהרות.'),
    ('יכולת לזכות בציון האיכות (40 נק\')', 0.15, 3, 'ניסיון מול רשויות, ראיון, 5 המלצות.'),
    ('תחרות וספק מקומי (העדפה של 5%)', 0.10, 2, 'ספק מקומי עשוי לזכות גם עם ציון נמוך יותר.'),
    ('משאבים ולוח זמנים', 0.05, 3, 'הגשה 28.10.2026.'),
]):
    r = 4 + i
    put(ws, f'A{r}', name).alignment = WRAP
    put(ws, f'B{r}', w, BLUE, PCT)
    put(ws, f'C{r}', s, BLUE, '0', YEL)
    put(ws, f'D{r}', f'=B{r}*C{r}', fmt='0.00')
    put(ws, f'E{r}', note).alignment = WRAP
    ws.row_dimensions[r].height = 30
put(ws, 'A12', 'סה"כ (מתוך 5)', BOLD)
put(ws, 'B12', '=SUM(B4:B10)', BOLD, PCT)
put(ws, 'D12', '=SUM(D4:D10)', BOLD, '0.00')
put(ws, 'E12', '=IF(C4<3,"לא להגיש – תנאי סף",IF(D12>=3.5,"להגיש",IF(D12>=2.75,"להגיש בתנאים – לפי תשובות ההבהרה","לא להגיש")))', BOLD)
put(ws, 'A13', 'רווח צפוי ל-36 חודשים (מקושר)')
put(ws, 'D13', "='תזרים'!C2", GREEN, ILS)
put(ws, 'A14', 'זוכה צפוי לפי ההנחות (מקושר)')
put(ws, 'D14', "=IF('ניקוד ותחרות'!N5<>\"\",\"HOT\",\"מתחרה\")", GREEN)

wb.save(OUT)
print('saved', OUT)
