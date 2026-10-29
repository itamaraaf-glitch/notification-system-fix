"""מודל כלכלי – נתיבי הגז הטבעי לישראל, מכרז INGL/TENDER/2026/48: נתבים ומתגים Cisco.

כתב הכמויות: 9 פריטי רכיבים (בדולר) + 7 שירותים (בשקלים). ניקוד: מחיר 60 (על הסך הכולל), איכות 40.
מחירי המחירון (GPL) ועלויות העבודה הם הערכות – לאמת מול מפיץ Cisco והמחלקות (גיליון "הנחות").

הרצה: python3 model.py <out.xlsx>   ואז חישוב מחדש (Excel / LibreOffice).
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule

OUT = sys.argv[1]
F = 'Arial'
BLUE, BLACK, GREEN = Font(name=F, color='0000FF'), Font(name=F), Font(name=F, color='008000')
BOLD, TITLE = Font(name=F, bold=True), Font(name=F, bold=True, size=14)
HDR, HFILL = Font(name=F, bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F4E79')
YEL = PatternFill('solid', fgColor='FFFF00')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
ILS, USD = '#,##0 ₪;(#,##0 ₪);-', '$#,##0;($#,##0);-'
NUM, PCT, DEC = '#,##0;(#,##0);-', '0.0%;(0.0%);-', '0.00'
WRAP = Alignment(wrap_text=True, vertical='top', horizontal='right')
RED = dict(font=Font(name=F, color='FFFFFF', bold=True), fill=PatternFill('solid', fgColor='C00000'))

wb = Workbook()


def sheet(name, widths):
    ws = wb.create_sheet(name)
    ws.sheet_view.rightToLeft = True
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    return ws


def header(ws, row, labels):
    for i, h in enumerate(labels):
        c = ws.cell(row=row, column=1 + i, value=h)
        c.font, c.fill, c.border = HDR, HFILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[row].height = 45


def put(ws, ref, val, font=BLACK, fmt=None, fill=None, note=None):
    c = ws[ref]
    c.value, c.font, c.border = val, font, BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if note:
        c.comment = Comment(note, 'מודל')
    return c


H = "'הנחות'!"
K = "'כתב כמויות'!"
N = "'ניקוד ותחרות'!"

# ------------------------------------------------------------------ מקרא
ws = wb.active
ws.title = 'מקרא'
ws.sheet_view.rightToLeft = True
ws.column_dimensions['A'].width = 120
for i, (t, f) in enumerate([
    ('מודל כלכלי – נתג"ז INGL/TENDER/2026/48: אספקה, התקנה ותחזוקה של נתבים ומתגים Cisco', TITLE),
    ('', BLACK),
    ('ניקוד: מחיר 60 (60 × הזולה / הנבדקת, על סך כל הפריטים; רכיבים בדולר מומרים בשער היציג ביום ההגשה), איכות 40 (מוסף 1). אין ערבויות. תשלום שוטף + 45. אחריות ותחזוקה 60 חודשים – אין לה סעיף בכתב הכמויות, לכן היא נטענת על מחירי הפריטים.', BOLD),
    ('', BLACK),
    ('מוסכמות: כחול = קלט. צהוב = הנחה לאמת. שחור = נוסחה. ירוק = קישור. אדום = הפסד.', BLACK),
    ('מחירי המחירון (GPL) בגיליון "כתב כמויות" הם הערכה גסה לחבילת המקטים של כל פריט – להחליף בהצעת מפיץ Cisco (Deal Registration).', BLACK),
    ('', BLACK),
    ('גיליונות:', BOLD),
    ('הנחות – שער דולר, הנחת קנייה מהמפיץ, הנחה מוצעת לנתג"ז, עלויות עבודה, רזרבת אחזקה.', BLACK),
    ('כתב כמויות – 16 הפריטים: כמות, מחיר מחירון, עלות, מחיר מוצע, מרווח.', BLACK),
    ('ניקוד ותחרות – ציון מחיר ואיכות מול מתחרה, מחיר מרבי לשוויון, שווי נקודה, טבלת רגישות להנחה.', BLACK),
    ('רווחיות – הכנסה, עלויות (כולל 60 חודשי אחזקה), רווח.', BLACK),
    ('תזרים – 66 חודשים: רכישה מראש, תקבולים בשוטף + 45, עלות אחזקה חודשית.', BLACK),
], start=1):
    c = ws.cell(row=i, column=1, value=t)
    c.font, c.alignment = f, Alignment(wrap_text=True, horizontal='right')
ws['A5'].fill = YEL

# ------------------------------------------------------------------ הנחות
ws = sheet('הנחות', {'A': 58, 'B': 14, 'C': 72})
ws['A1'] = 'הנחות'
ws['A1'].font = TITLE
header(ws, 3, ['הנחה', 'ערך', 'הסבר'])
rows = [
    (4, 'שער דולר יציג (₪)', 3.65, DEC, 'השוואת ההצעות לפי השער היציג במועד האחרון להגשה (13.12.2026). התשלום – לפי השער ביום ההזמנה.'),
    (5, 'הנחת קנייה מהמפיץ (מ-GPL)', 0.45, PCT, 'הערכה – תלויה ברישום העסקה (Deal Registration) ובדרגת השותף. לאמת מול המפיץ.'),
    (6, 'הנחה מוצעת לנתג"ז (מ-GPL)', 0.35, PCT, 'הקלט המרכזי לתחרות. כל המתחרים קונים מאותם מפיצים – הפער ביניהם הוא בעיקר ברישום העסקה.'),
    (7, 'תוספת רווח על עלות השירותים', 0.30, PCT, ''),
    (8, 'עלות יום טכנאי שטח (₪)', 1800, ILS, 'כולל רכב ונסיעות לאתרים מדימונה עד חיפה.'),
    (9, 'עלות יום מהנדס רשת / אבטחה (₪)', 2600, ILS, ''),
    (10, 'עלות יום מהנדס CCIE (₪)', 4000, ILS, 'ייתכן קבלן משנה – שאלת הבהרה 21.'),
    (11, 'עלות חודשית לאחזקה 24/7 בשטח (₪)', 3000, ILS, 'אין סעיף אחזקה בכתב הכמויות: כוננות, הגעה לאתרים, טיפול בתקלות – 60 חודשים. מקטי תמיכת היצרן כלולים במחיר הרכיבים.'),
    (12, 'חודשי אחזקה', 60, '0', 'מפרט 2.3.'),
    (13, 'מרווח ביטחון על עלות השירותים', 0.10, PCT, 'גישה לאתרים, היתרים, ליווי, עבודה בגובה.'),
    (14, 'פיגור תשלום (חודשים)', 2.5, '0.0', 'שוטף + 45 מקבלת החשבון + אישור הממונה.'),
    (15, 'שיעור היוון שנתי', 0.08, PCT, ''),
    (16, 'חודש אספקת הציוד', 4, '0', 'הערכה – זמני אספקה של Cisco (שאלת הבהרה 5).'),
    (17, 'משך ההתקנות (חודשים)', 6, '0', '60 מתקנים, עד 3 ביום, בתיאום עם החברה.'),
]
for r, lab, val, fmt, note in rows:
    put(ws, f'A{r}', lab).alignment = WRAP
    put(ws, f'B{r}', val, BLUE, fmt, YEL)
    put(ws, f'C{r}', note).alignment = WRAP
    ws.row_dimensions[r].height = 30

# ------------------------------------------------------------------ כתב כמויות
ws = sheet('כתב כמויות', {'A': 5, 'B': 44, 'C': 9, 'D': 9, 'E': 13, 'F': 13, 'G': 13, 'H': 15, 'I': 15, 'J': 10, 'K': 46})
ws['A1'] = 'כתב כמויות – מחיר מחירון, עלות ומחיר מוצע'
ws['A1'].font = TITLE
ws['A2'] = 'רכיבים: E = מחיר מחירון (GPL) ליחידה בדולר. שירותים: E = ימי עבודה ליחידה; F/G בשקלים.'
header(ws, 3, ['#', 'פריט', 'סוג', 'כמות', 'GPL $ / ימים ליחידה', 'עלות ליחידה', 'מחיר מוצע ליחידה',
               'סה"כ מחיר (₪)', 'סה"כ עלות (₪)', 'מרווח', 'הערה (הערכה – לאמת)'])
comps = [
    (1, 'מתג IE3500 8 פורטים + Network Advantage + SD + תמיכה', 20, 3500, 'חבילת 9 מקטים.'),
    (2, 'מודול הרחבה IEM-3500-16T + תמיכה', 20, 1800, ''),
    (3, 'נתב IR1101-K9 + Network Essentials + אנטנות + קיט', 60, 3000, 'חבילת 16 מקטים.'),
    (4, 'יחידת 5G RedCap (P-5GS6-RC-GL)', 60, 1200, ''),
    (5, 'הרחבת SSD לאחסון לוגים (IRM-1100-SPMI + 100G)', 10, 900, ''),
    (6, 'ISE – רישוי מכונה וירטואלית (R-ISE-VMC-K9)', 3, 600, 'אין חומרה בכתב הכמויות – שאלת הבהרה 19.'),
    (7, 'ISE Advantage (מנוי, לנקודת קצה)', 1440, 90, 'משך המנוי לא מוגדר – שאלת הבהרה 15.'),
    (8, 'ISE PLR – רישוי לסביבה מנותקת', 3, 500, ''),
    (9, 'Cyber Vision Advantage, 5 שנים (לנכס)', 1440, 120, 'אין פריט ל-Cyber Vision Center – שאלת הבהרה 20.'),
]
R0 = 4
for i, (n, name, qty, gpl, note) in enumerate(comps):
    r = R0 + i
    put(ws, f'A{r}', n)
    put(ws, f'B{r}', name).alignment = WRAP
    put(ws, f'C{r}', 'רכיב')
    put(ws, f'D{r}', qty, BLUE, NUM)
    put(ws, f'E{r}', gpl, BLUE, USD, YEL)
    put(ws, f'F{r}', f'=E{r}*(1-{H}$B$5)', fmt=USD)
    put(ws, f'G{r}', f'=E{r}*(1-{H}$B$6)', fmt=USD)
    put(ws, f'H{r}', f'=D{r}*G{r}*{H}$B$4', fmt=ILS)
    put(ws, f'I{r}', f'=D{r}*F{r}*{H}$B$4', fmt=ILS)
    put(ws, f'J{r}', f'=IF(H{r}>0,(H{r}-I{r})/H{r},0)', fmt=PCT)
    put(ws, f'K{r}', note).alignment = WRAP
    ws.row_dimensions[r].height = 30
RC = R0 + len(comps) - 1  # 12
# שירותים: (n, name, qty, days, rate_cell, note)
servs = [
    (10, 'התקנת נתב במתקן (מעבדה + שטח)', 60, 0.67, 'B8', 'טכנאי שטח + טכנאי באתר הראשי, 3 התקנות ביום (מפרט 3.2.2).'),
    (11, 'התקנת 2 אנטנות חיצוניות במתקן', 60, 0.6, 'B8', 'עבודה בגובה, 2 × 15 מ\' קואקס; חומרים בנפרד – ראו הערה.'),
    (12, 'התקנה והטמעת ISE (Cluster, Profiling, MAB, Micro-segmentation)', 1, 18, 'B9', ''),
    (13, 'התקנה והטמעת Cyber Vision', 1, 10, 'B9', ''),
    (14, 'התקנת 2 מתגים במתקן + חיבור Sensor', 10, 0.8, 'B9', 'איש רשת + איש אבטחת מידע (מפרט 3.4.2).'),
    (15, 'מסמכי HLD + LLD + תיקי אתר (CCIE)', 1, 12, 'B10', ''),
    (16, 'ניהול פרויקט (קומפלט)', 1, 25, 'B9', 'כ-25% משרה לאורך כ-6 חודשים.'),
]
RS0 = RC + 1
for i, (n, name, qty, days, rate, note) in enumerate(servs):
    r = RS0 + i
    put(ws, f'A{r}', n)
    put(ws, f'B{r}', name).alignment = WRAP
    put(ws, f'C{r}', 'שירות')
    put(ws, f'D{r}', qty, BLUE, NUM)
    put(ws, f'E{r}', days, BLUE, '0.00', YEL)
    extra = '+250' if n == 11 else ''
    put(ws, f'F{r}', f'=E{r}*{H}${rate[0]}${rate[1:]}*(1+{H}$B$13){extra}', fmt=ILS,
        note='כולל 250 ₪ חומרים (קואקס, תושבות) למתקן.' if n == 11 else None)
    put(ws, f'G{r}', f'=F{r}*(1+{H}$B$7)', fmt=ILS)
    put(ws, f'H{r}', f'=D{r}*G{r}', fmt=ILS)
    put(ws, f'I{r}', f'=D{r}*F{r}', fmt=ILS)
    put(ws, f'J{r}', f'=IF(H{r}>0,(H{r}-I{r})/H{r},0)', fmt=PCT)
    put(ws, f'K{r}', note).alignment = WRAP
    ws.row_dimensions[r].height = 30
RL = RS0 + len(servs) - 1  # 19
r = RL + 2
put(ws, f'B{r}', 'סה"כ רכיבים', BOLD)
put(ws, f'H{r}', f'=SUM(H{R0}:H{RC})', BOLD, ILS)
put(ws, f'I{r}', f'=SUM(I{R0}:I{RC})', BOLD, ILS)
put(ws, f'B{r+1}', 'סה"כ שירותים', BOLD)
put(ws, f'H{r+1}', f'=SUM(H{RS0}:H{RL})', BOLD, ILS)
put(ws, f'I{r+1}', f'=SUM(I{RS0}:I{RL})', BOLD, ILS)
put(ws, f'B{r+2}', 'סה"כ הצעה (₪) – הבסיס לציון המחיר', BOLD)
put(ws, f'H{r+2}', f'=H{r}+H{r+1}', BOLD, ILS)
put(ws, f'I{r+2}', f'=I{r}+I{r+1}', BOLD, ILS)
put(ws, f'B{r+3}', 'נתח הרכיבים מסך ההצעה')
put(ws, f'H{r+3}', f'=IF(H{r+2}>0,H{r}/H{r+2},0)', fmt=PCT)
T_COMP, T_SERV, T_ALL = f'{K}$H${r}', f'{K}$H${r+1}', f'{K}$H${r+2}'
C_COMP, C_SERV = f'{K}$I${r}', f'{K}$I${r+1}'
ws.conditional_formatting.add(f'J{R0}:J{RL}', CellIsRule(operator='lessThan', formula=['0'], **RED))
ws.freeze_panes = 'C4'

# ------------------------------------------------------------------ ניקוד ותחרות
ws = sheet('ניקוד ותחרות', {'A': 52, 'B': 15, 'C': 15, 'D': 15, 'E': 50})
ws['A1'] = 'ניקוד ותחרות'
ws['A1'].font = TITLE
header(ws, 3, ['רכיב', 'מרבי', 'HOT', 'מתחרה', 'הערה'])
put(ws, 'A4', 'הנחת המתחרה מ-GPL (רכיבים)')
put(ws, 'D4', 0.40, BLUE, PCT, YEL)
put(ws, 'E4', 'שותף עם רישוי עסקה (Deal Registration) מקבל הנחה גבוהה יותר.').alignment = WRAP
put(ws, 'A5', 'מחיר השירותים של המתחרה ביחס לשלנו')
put(ws, 'D5', 0.90, BLUE, PCT, YEL)
put(ws, 'A6', 'סך ההצעה (₪)', BOLD)
put(ws, 'C6', f'={T_ALL}', BOLD, ILS)
put(ws, 'D6', f"=SUMPRODUCT({K}D{R0}:D{RC},{K}E{R0}:E{RC})*(1-D4)*{H}B4+{T_SERV}*D5", BOLD, ILS)
put(ws, 'A7', 'ציון מחיר')
put(ws, 'B7', 60, fmt=NUM)
put(ws, 'C7', '=60*MIN(C6,D6)/C6', fmt=DEC)
put(ws, 'D7', '=60*MIN(C6,D6)/D6', fmt=DEC)
qual = [
    (8, 'הסמכת Preferred Security Partner (א)', 10, 0, 10, 'לא ידוע אם ל-HOT יש – שאלה פנימית.'),
    (9, 'פרויקטי Cyber Vision מעבר לראשון (ב) – 2.5 לפרויקט', 10, 0, 5, ''),
    (10, 'עובדי CCNA (ג) – 11–20: 5; 21–30: 7.5; 31+: 10', 10, 5, 10, ''),
    (11, 'עובדי CCST (ג) – 3–5: 2.5; 6–8: 5; 9–11: 7.5; 12+: 10', 10, 2.5, 5, ''),
]
for r, lab, mx, ours, comp, note in qual:
    put(ws, f'A{r}', lab).alignment = WRAP
    put(ws, f'B{r}', mx, fmt=NUM)
    put(ws, f'C{r}', ours, BLUE, DEC, YEL)
    put(ws, f'D{r}', comp, BLUE, DEC, YEL)
    put(ws, f'E{r}', note).alignment = WRAP
put(ws, 'A12', 'ציון כולל', BOLD)
put(ws, 'B12', 100, BOLD, NUM)
put(ws, 'C12', '=SUM(C7:C11)', BOLD, DEC)
put(ws, 'D12', '=SUM(D7:D11)', BOLD, DEC)
put(ws, 'A13', 'פער (HOT פחות מתחרה)', BOLD)
put(ws, 'C13', '=C12-D12', BOLD, '0.00;[Red]-0.00')
put(ws, 'A14', 'הצעה מרבית של HOT לשוויון ציונים (₪)').alignment = WRAP
put(ws, 'C14', '=60*D6/(60+SUM(D8:D11)-SUM(C8:C11))', BOLD, ILS)
put(ws, 'E14', 'אם הערך נמוך מהעלות (גיליון רווחיות) – אין דרך לנצח במחיר בלבד.').alignment = WRAP
put(ws, 'A15', 'שווי נקודת איכות אחת בהצעה (₪)')
put(ws, 'C15', '=C6*C6/(60*MIN(C6,D6))', BOLD, ILS)
put(ws, 'E15', 'כמה ₪ צריך להוריד מההצעה כדי לפצות על נקודה אחת.').alignment = WRAP

put(ws, 'A17', 'רגישות: הנחה מוצעת מ-GPL', BOLD)
header(ws, 18, ['הנחה מוצעת', 'סך הצעה (₪)', 'ציון מחיר', 'ציון כולל HOT', 'רווח לפני אחזקה (₪)'])
for i, d in enumerate([0.25, 0.30, 0.35, 0.40, 0.42, 0.45]):
    r = 19 + i
    put(ws, f'A{r}', d, BLUE, PCT)
    put(ws, f'B{r}', f"=SUMPRODUCT({K}$D${R0}:$D${RC},{K}$E${R0}:$E${RC})*(1-A{r})*{H}$B$4+{T_SERV}", fmt=ILS)
    put(ws, f'C{r}', f'=60*MIN(B{r},$D$6)/B{r}', fmt=DEC)
    put(ws, f'D{r}', f'=C{r}+SUM($C$8:$C$11)', fmt=DEC)
    put(ws, f'E{r}', f'=B{r}-{C_COMP}-{C_SERV}', fmt=ILS)
ws.conditional_formatting.add('E19:E24', CellIsRule(operator='lessThan', formula=['0'], **RED))

# ------------------------------------------------------------------ רווחיות
ws = sheet('רווחיות', {'A': 50, 'B': 16, 'C': 50})
ws['A1'] = 'רווחיות לכל ההתקשרות'
ws['A1'].font = TITLE
header(ws, 3, ['סעיף', '₪', 'הערה'])
prof = [
    (4, 'הכנסה – רכיבים', f'={T_COMP}', ''),
    (5, 'הכנסה – שירותים', f'={T_SERV}', ''),
    (6, 'סה"כ הכנסה', '=B4+B5', ''),
    (7, 'עלות רכיבים', f'=-{C_COMP}', 'בדולר – חשיפה לשער עד מועד הרכישה.'),
    (8, 'עלות שירותים', f'=-{C_SERV}', ''),
    (9, 'אחזקה 24/7 ל-60 חודשים', f"=-{H}B11*{H}B12", 'אין עליה תמורה נפרדת.'),
    (10, 'סה"כ עלות', '=B7+B8+B9', ''),
    (11, 'רווח', '=B6+B10', ''),
    (12, 'מרווח', '=IF(B6>0,B11/B6,0)', ''),
]
for r, lab, f, note in prof:
    put(ws, f'A{r}', lab, BOLD if r in (6, 10, 11, 12) else BLACK)
    put(ws, f'B{r}', f, BOLD if r in (6, 10, 11, 12) else BLACK, PCT if r == 12 else ILS)
    put(ws, f'C{r}', note).alignment = WRAP
ws.conditional_formatting.add('B11:B12', CellIsRule(operator='lessThan', formula=['0'], **RED))

# ------------------------------------------------------------------ תזרים
ws = sheet('תזרים', {c: 14 for c in 'ABCDEFG'})
ws.column_dimensions['A'].width = 8
ws['A1'] = 'תזרים חודשי (66 חודשים)'
ws['A1'].font = TITLE
R1, LAST = 7, 7 + 66
put(ws, 'A2', 'NPV', BOLD)
ws.merge_cells('A2:B2')
put(ws, 'C2', f"=F{R1}+NPV({H}B15/12,F{R1+1}:F{LAST})", BOLD, ILS)
put(ws, 'A3', 'חשיפת מזומנים מרבית', BOLD)
ws.merge_cells('A3:B3')
put(ws, 'C3', f'=MIN(G{R1}:G{LAST})', BOLD, ILS)
put(ws, 'A4', 'הנחה: הציוד משולם למפיץ חודש לפני האספקה; חשבון רכיבים באספקה; שירותים בחלקים שווים לאורך ההתקנות; תקבול אחרי פיגור התשלום; אחזקה מחודש האספקה.')
ws.merge_cells('A4:G4')
header(ws, 6, ['חודש', 'חשבון רכיבים', 'חשבון שירותים', 'תקבול', 'תשלום לספקים ועבודה', 'תזרים נטו', 'מצטבר'])
DM, IN = f'{H}$B$16', f'{H}$B$17'
LAG = f'ROUND({H}$B$14,0)'
for m in range(0, 67):
    r = R1 + m
    put(ws, f'A{r}', m, fmt='0')
    put(ws, f'B{r}', f'=IF(A{r}={DM},{T_COMP},0)', fmt=ILS)
    put(ws, f'C{r}', f'=IF(AND(A{r}>={DM},A{r}<{DM}+{IN}),{T_SERV}/{IN},0)', fmt=ILS)
    put(ws, f'D{r}', f'=IF(A{r}-{LAG}>=0,INDEX($B${R1}:$B${LAST},A{r}-{LAG}+1)+INDEX($C${R1}:$C${LAST},A{r}-{LAG}+1),0)', fmt=ILS)
    put(ws, f'E{r}', f'=-(IF(A{r}={DM}-1,{C_COMP},0)+IF(AND(A{r}>={DM},A{r}<{DM}+{IN}),{C_SERV}/{IN},0)'
                     f'+IF(AND(A{r}>={DM},A{r}<{DM}+{H}$B$12),{H}$B$11,0))', fmt=ILS)
    put(ws, f'F{r}', f'=D{r}+E{r}', fmt=ILS)
    put(ws, f'G{r}', f'=F{r}' if m == 0 else f'=G{r-1}+F{r}', fmt=ILS)
ws.freeze_panes = f'B{R1}'

wb.save(OUT)
print('saved', OUT)
