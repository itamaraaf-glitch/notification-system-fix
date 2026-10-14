"""מודל כלכלי – מכרז עירייה (שם לא ידוע): קווי תקשורת, Fortinet, מרכזייה בענן.
מבוסס על כתב הכמויות בלבד (חוברת המכרז לא התקבלה).

הרצה: python3 model.py <out.xlsx>   ואז חישוב מחדש (Excel / LibreOffice).
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule

OUT = sys.argv[1]
F = 'Arial'
BLUE, BLACK, GREEN = Font(name=F, color='0000FF'), Font(name=F), Font(name=F, color='008000')
BOLD, TITLE = Font(name=F, bold=True), Font(name=F, bold=True, size=14)
HDR, HFILL = Font(name=F, bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F4E79')
YEL = PatternFill('solid', fgColor='FFFF00')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
ILS, ILS2 = '#,##0 ₪;(#,##0 ₪);-', '#,##0.00 ₪;(#,##0.00 ₪);-'
NUM, PCT = '#,##0;(#,##0);-', '0.0%;(0.0%);-'
WRAP = Alignment(wrap_text=True, vertical='top', horizontal='right')
RED = dict(font=Font(name=F, color='FFFFFF', bold=True), fill=PatternFill('solid', fgColor='C00000'))

wb = Workbook()


def sheet(name, widths):
    ws = wb.create_sheet(name)
    ws.sheet_view.rightToLeft = True
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    return ws


def header(ws, row, labels):
    for i, h in enumerate(labels):
        c = ws.cell(row=row, column=1 + i, value=h)
        c.font, c.fill, c.border = HDR, HFILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[row].height = 48


def put(ws, ref, val, font=BLACK, fmt=None, fill=None, note=None):
    c = ws[ref]
    c.value, c.font, c.border = val, font, BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if note:
        c.comment = Comment(note, 'מודל')
    return c


H = "'הנחות'!"
K = "'כתב כמויות'!"

# ------------------------------------------------------------------ מקרא
ws = wb.active
ws.title = 'מקרא'
ws.sheet_view.rightToLeft = True
ws.column_dimensions['A'].width = 115
for i, (t, f) in enumerate([
    ('מודל כלכלי – מכרז עירייה: קווי תקשורת, Fortinet ומרכזייה בענן (שם המכרז לא ידוע)', TITLE),
    ('', BLACK),
    ('המודל מבוסס על כתב הכמויות בלבד. חסרים מחוברת המכרז: שיטת הניקוד, תקופת ההתקשרות, תנאי תשלום, הצמדה ותנאי סף. כשהחוברת תגיע – לעדכן את גיליון "הנחות".', BOLD),
    ('', BLACK),
    ('מוסכמות: כחול = קלט. צהוב = הנחה לאמת מול HOT. שחור = נוסחה. ירוק = קישור מגיליון אחר. אדום = רכיב מפסיד או שגיאה בטופס.', BLACK),
    ('', BLACK),
    ('גיליונות:', BOLD),
    ('הנחות – תקופה, הנחה לכל קטגוריה, פיגור תשלום, היוון, האם הספק המכהן מגיש.', BLACK),
    ('כתב כמויות – 29 הסעיפים, מחיר מרבי, ההצעה שלנו, עלות, מרווח, וחלוקה לחודשי / חד-פעמי.', BLACK),
    ('סיכום והשוואה – סך ההצעה לפי שני בסיסי השוואה אפשריים, ורווחיות לפי קטגוריה.', BLACK),
    ('תזרים – 60 חודשים: חומרה ועבודה מראש, הכנסה חודשית, NPV וחשיפה מרבית.', BLACK),
    ('שגיאות בטופס – ליקויים בכתב הכמויות שמשפיעים על ההשוואה (בסיס לשאלות ההבהרה).', BLACK),
], start=1):
    c = ws.cell(row=i, column=1, value=t)
    c.font, c.alignment = f, Alignment(wrap_text=True, horizontal='right')
ws['A5'].fill = YEL

# ------------------------------------------------------------------ הנחות
ws = sheet('הנחות', {'A': 56, 'B': 14, 'C': 70})
ws['A1'] = 'הנחות'
ws['A1'].font = TITLE
header(ws, 3, ['הנחה', 'ערך', 'הסבר'])
rows = [
    (4, 'תקופת ההתקשרות הבסיסית (חודשים)', 36, '0', 'לא ידוע – הערות הטופס מזכירות תמיכה ל-3 שנים וסעיף 9 "החל משנה רביעית".'),
    (5, 'חודשי אופציה צפויים (מעבר לבסיס)', 0, '0', 'לא ידוע. סעיף 9 (תחזוקה 2,000 ₪) מתחיל רק בשנה הרביעית.'),
    (6, 'הנחה על קווי תקשורת (סעיפים 10–22)', 0.0, PCT, 'ליבת HOT – המקום לתחרות.'),
    (7, 'הנחה על מרכזייה בענן (סעיף 4)', 0.0, PCT, ''),
    (8, 'הנחה על רישוי Fortinet SaaS (סעיפים 1–3)', 0.0, PCT, 'מחיר יצרן – מרווח מוגבל.'),
    (9, 'הנחה על חומרה (סעיפים 23–28)', 0.0, PCT, 'מחיר יצרן צמוד דולר.'),
    (10, 'הנחה על שעות עבודה (סעיפים 5–8, 29)', 0.0, PCT, ''),
    (11, 'הנחה על תחזוקה משנה רביעית (סעיף 9)', 0.0, PCT, ''),
    (12, 'הספק המכהן מגיש הצעה? (1 = כן, 0 = לא)', 1, '0', 'לפי הערות הטופס: אם כן – סעיף 7 לא נכלל בהשוואה, ומכסת השעות שלו יורדת ל-10.'),
    (13, 'פיגור תשלום (חודשים)', 2, '0', 'לא ידוע – הנחה (שוטף + 45 ועוד אישור).'),
    (14, 'שיעור היוון שנתי', 0.08, PCT, ''),
    (15, 'חודש התקנת ציוד וקווים (תחילת חיוב חודשי)', 2, '0', 'הנחה.'),
    (16, 'עלות שעת עבודה פנימית', 160, ILS, 'הערכה – לאמת.'),
]
for r, lab, val, fmt, note in rows:
    put(ws, f'A{r}', lab).alignment = WRAP
    put(ws, f'B{r}', val, BLUE, fmt, YEL)
    put(ws, f'C{r}', note).alignment = WRAP
put(ws, 'A18', 'נגזרות', BOLD)
put(ws, 'A19', 'סה"כ חודשי התקשרות')
put(ws, 'B19', '=B4+B5', fmt='0')
put(ws, 'A20', 'חודשים עם תחזוקה משנה רביעית (סעיף 9)')
put(ws, 'B20', '=MAX(0,B19-36)', fmt='0')

# ------------------------------------------------------------------ כתב כמויות
ws = sheet('כתב כמויות', {'A': 5, 'B': 46, 'C': 12, 'D': 13, 'E': 8, 'F': 12, 'G': 12, 'H': 12, 'I': 13, 'J': 13, 'K': 13, 'L': 10, 'M': 44})
ws['A1'] = 'כתב כמויות – מחירים מרביים, ההצעה שלנו ועלויות'
ws['A1'].font = TITLE
header(ws, 3, ['#', 'רכיב', 'קטגוריה', 'סוג תשלום', 'כמות', 'מחיר מרבי ליחידה', 'המחיר שלנו ליחידה', 'עלות שלנו ליחידה',
               'סה"כ מרבי', 'סה"כ שלנו', 'סה"כ עלות', 'מרווח', 'הערת עלות (הערכה – לאמת)'])
# (n, name, category, pay, qty, max, cost, note)  pay: חודשי / חד-פעמי / חודשי משנה 4
L, P, S, HW, LB, MT = 'קווים', 'מרכזייה', 'Fortinet SaaS', 'חומרה', 'שעות עבודה', 'תחזוקה'
DISC = {L: 'B6', P: 'B7', S: 'B8', HW: 'B9', LB: 'B10', MT: 'B11'}
items = [
    (1, 'FortiNAC Cloud SaaS ל-400 התקנים', S, 'חודשי', 1, 1700, 1300, '~75% ממחיר מרבי.'),
    (2, 'FortiAnalyzer Cloud SaaS (NTA/NPM), שמירה 24 ח\', FortiCare', S, 'חודשי', 1, 1900, 1450, ''),
    (3, 'FortiManager Cloud SaaS (NCM) עד 50 התקנים, FortiCare', S, 'חודשי', 1, 1500, 1150, ''),
    (4, 'מרכזייה בענן + DR, 200 שלוחות, SIP trunk, טווח מספרים, Call Center, 10,000 דקות, 2037*', P, 'חודשי', 1, 3500, 1900, 'פלטפורמת HOT: 200 × ~6 ₪ + דקות + DR. עלות 2037* לא ידועה.'),
    (5, 'הקמת מערך FortiNAC (מכסת שעות)', LB, 'חד-פעמי', 50, 250, "='הנחות'!B16", ''),
    (6, 'הקמת מערך ניהול רשת (FortiAnalyzer & FortiManager)', LB, 'חד-פעמי', 60, 250, "='הנחות'!B16", ''),
    (7, 'הקמת מרכזייה Hosted IPT (מכסת שעות)', LB, 'חד-פעמי', 50, 220, "='הנחות'!B16", 'לא נכלל בהשוואה אם המכהן מגיש.'),
    (8, 'הגדרת FortiGate, מתגים ורשתות אלחוטיות', LB, 'חד-פעמי', 85, 220, "='הנחות'!B16", ''),
    (9, 'תחזוקת רשת פנימית (מתגים, AP, מחשבים) – משנה רביעית', MT, 'חודשי משנה 4', 1, 2000, 1400, ''),
    (10, 'קו IP/מטרו 200M + אינטרנט + IP', L, 'חודשי', 1, 1850, 750, 'עלות On-Net – הערכה.'),
    (11, 'קו IP/מטרו 150M + אינטרנט + IP', L, 'חודשי', 6, 1400, 600, ''),
    (12, 'קו IP/מטרו 300/150M + אינטרנט + IP', L, 'חודשי', 1, 1600, 650, 'זול מ-200M (סעיף 10) – חריג.'),
    (13, 'קו אינטרנט על סיב 500/100M + IP (IPVPN)', L, 'חודשי', 24, 240, 120, ''),
    (14, 'קו IPVPN GPON 500/500M + ניהול ציוד + חיבור טלפון', L, 'חודשי', 24, 300, 170, 'בטופס המקורי הסעיף לא נסכם!'),
    (15, 'קו P2P 150M', L, 'חודשי', 1, 1400, 600, ''),
    (16, 'קו P2P 200M', L, 'חודשי', 1, 1850, 700, ''),
    (17, 'קו תמסורת ייעודי 1G לאתר ראשי + ניטור', L, 'חודשי', 1, 2000, 900, ''),
    (18, 'קו Private Profile לסניפים + ניטור', L, 'חודשי', 1, 400, 200, ''),
    (19, 'ציוד תקשורת לאתר ראשי (שכירות)', L, 'חודשי', 1, 300, 150, ''),
    (20, 'התקנה חד-פעמית לאתר (סעיפים 17+18)', L, 'חד-פעמי', 1, 1500, 1200, 'בטופס המקורי מופיע בעמודה החודשית!'),
    (21, 'קו IPVPN GPON 1000/1000M + ניהול ציוד + חיבור טלפון', L, 'חודשי', 1, 560, 300, ''),
    (22, 'קו אופטי GPON 1000/1000M ללא אינטרנט', L, 'חודשי', 0, 400, 220, 'כמות לא צוינה בטופס.'),
    (23, 'מתג FS-124F-FPOE + 36 ח\' תמיכה (או שו"ע)', HW, 'חד-פעמי', 15, 4700, 4200, 'מחיר מפיץ – לאמת. צמוד דולר.'),
    (24, 'מתג FS-148F-FPOE + 36 ח\' תמיכה (או שו"ע)', HW, 'חד-פעמי', 1, 7000, 6300, ''),
    (25, 'FortiGate FG-71G + 36 ח\' FortiCare (או שו"ע)', HW, 'חד-פעמי', 4, 6500, 5800, ''),
    (26, 'FortiGate FG-121G + 36 ח\' FortiCare (או שו"ע)', HW, 'חד-פעמי', 1, 23000, 20500, ''),
    (27, 'FortiGate FG-51G + 36 ח\' FortiCare (או שו"ע)', HW, 'חד-פעמי', 24, 4500, 4000, ''),
    (28, 'FortiAP 231G + 3Y UTM + FortiCare', HW, 'חד-פעמי', 10, 1200, 1400, 'חשד: מחיר מרבי נמוך מעלות השוק – לאמת.'),
    (29, 'פירוק ציוד ישן', LB, 'חד-פעמי', 30, 150, "='הנחות'!B16", 'מחיר מרבי 150 ₪ – נמוך מעלות שעה.'),
]
R0 = 4
for i, (n, name, cat, pay, qty, mx, cost, note) in enumerate(items):
    r = R0 + i
    put(ws, f'A{r}', n)
    put(ws, f'B{r}', name).alignment = WRAP
    put(ws, f'C{r}', cat)
    put(ws, f'D{r}', pay)
    put(ws, f'E{r}', qty, BLUE, NUM, YEL if n == 22 else None)
    put(ws, f'F{r}', mx, BLUE, ILS)
    put(ws, f'G{r}', f"=F{r}*(1-{H}{DISC[cat]})", fmt=ILS2)
    put(ws, f'H{r}', cost, GREEN if str(cost).startswith('=') else BLUE, ILS2, YEL)
    put(ws, f'I{r}', f'=E{r}*F{r}', fmt=ILS)
    put(ws, f'J{r}', f'=E{r}*G{r}', fmt=ILS)
    put(ws, f'K{r}', f'=E{r}*H{r}', fmt=ILS)
    put(ws, f'L{r}', f'=IF(G{r}>0,(G{r}-H{r})/G{r},0)', fmt=PCT)
    put(ws, f'M{r}', note).alignment = WRAP
    ws.row_dimensions[r].height = 30
RL = R0 + len(items) - 1
ws.conditional_formatting.add(f'L{R0}:L{RL}', CellIsRule(operator='lessThan', formula=['0'], **RED))
ws.freeze_panes = 'C4'
rng = lambda c: f'${c}${R0}:${c}${RL}'

# ------------------------------------------------------------------ סיכום והשוואה
ws = sheet('סיכום והשוואה', {'A': 50, 'B': 15, 'C': 15, 'D': 15, 'E': 12})
ws['A1'] = 'סיכום והשוואה'
ws['A1'].font = TITLE
header(ws, 3, ['סכום', 'מרבי', 'שלנו', 'עלות', 'מרווח'])
lines = [
    ('חודשי (מהחודש הראשון)', 'חודשי'),
    ('חודשי משנה רביעית (סעיף 9)', 'חודשי משנה 4'),
    ('חד-פעמי', 'חד-פעמי'),
]
for i, (lab, pay) in enumerate(lines):
    r = 4 + i
    put(ws, f'A{r}', lab)
    for col, src in zip('BCD', 'IJK'):
        put(ws, f'{col}{r}', f'=SUMIFS({K}{rng(src)},{K}{rng("D")},"{pay}")', fmt=ILS)
    put(ws, f'E{r}', f'=IF(C{r}>0,(C{r}-D{r})/C{r},0)', fmt=PCT)
put(ws, 'A8', 'סעיף 7 (מוחרג מההשוואה אם המכהן מגיש)')
for col, src in zip('BCD', 'IJK'):
    put(ws, f'{col}8', f"=INDEX({K}{rng(src)},7)", fmt=ILS)
put(ws, 'A10', 'בסיס השוואה א\': סה"כ חודשי בלבד (כמו הנוסחאות בטופס)', BOLD).alignment = WRAP
put(ws, 'B10', '=B4', BOLD, ILS)
put(ws, 'C10', '=C4', BOLD, ILS)
put(ws, 'A11', 'בסיס השוואה ב\': עלות כוללת לתקופה (חודשי × חודשים + חד-פעמי)', BOLD).alignment = WRAP
put(ws, 'B11', f"=B4*{H}B19+B5*{H}B20+B6-{H}B12*B8", BOLD, ILS)
put(ws, 'C11', f"=C4*{H}B19+C5*{H}B20+C6-{H}B12*C8", BOLD, ILS)
put(ws, 'A12', 'הערה: שיטת ההשוואה לא מוגדרת בטופס (שורת הסה"כ ריקה). שאלת הבהרה 2.').alignment = WRAP
ws.merge_cells('A12:E12')

put(ws, 'A14', 'רווחיות לתקופה', BOLD)
header(ws, 15, ['קטגוריה', 'הכנסה לתקופה', 'עלות לתקופה', 'רווח', 'מרווח'])
cats = [L, P, S, HW, LB, MT]
for i, cat in enumerate(cats):
    r = 16 + i
    put(ws, f'A{r}', cat)
    mult = f"IF({K}{rng('D')}=\"חודשי\",{H}$B$19,IF({K}{rng('D')}=\"חודשי משנה 4\",{H}$B$20,1))"
    put(ws, f'B{r}', f'=SUMPRODUCT(({K}{rng("C")}="{cat}")*{K}{rng("J")}*{mult})', fmt=ILS)
    put(ws, f'C{r}', f'=SUMPRODUCT(({K}{rng("C")}="{cat}")*{K}{rng("K")}*{mult})', fmt=ILS)
    put(ws, f'D{r}', f'=B{r}-C{r}', fmt=ILS)
    put(ws, f'E{r}', f'=IF(B{r}>0,D{r}/B{r},0)', fmt=PCT)
put(ws, 'A22', 'סה"כ', BOLD)
for col in 'BCD':
    put(ws, f'{col}22', f'=SUM({col}16:{col}21)', BOLD, ILS)
put(ws, 'E22', '=IF(B22>0,D22/B22,0)', BOLD, PCT)
ws.conditional_formatting.add('E16:E22', CellIsRule(operator='lessThan', formula=['0'], **RED))

# ------------------------------------------------------------------ תזרים
ws = sheet('תזרים', {c: 14 for c in 'ABCDEFGH'})
ws.column_dimensions['A'].width = 8
ws['A1'] = 'תזרים חודשי (60 חודשים)'
ws['A1'].font = TITLE
R1, LAST = 8, 8 + 60
S_ = "'סיכום והשוואה'!"
put(ws, 'A2', 'רווח (צבירה) לתקופה', BOLD)
ws.merge_cells('A2:B2')
put(ws, 'C2', f'=SUM(E{R1}:E{LAST})', BOLD, ILS)
put(ws, 'A3', 'NPV של התזרים', BOLD)
ws.merge_cells('A3:B3')
put(ws, 'C3', f"=G{R1}+NPV({H}B14/12,G{R1+1}:G{LAST})", BOLD, ILS)
put(ws, 'A4', 'חשיפת מזומנים מרבית', BOLD)
ws.merge_cells('A4:B4')
put(ws, 'C4', f'=MIN(H{R1}:H{LAST})', BOLD, ILS)
header(ws, 7, ['חודש', 'הכנסה (צבירה)', 'עלות חודשית', 'עלות חד-פעמית', 'רווח (צבירה)', 'תקבול', 'תזרים נטו', 'מצטבר'])
GL = f'{H}$B$15'
for m in range(0, 61):
    r = R1 + m
    act = f'AND(A{r}>={GL},A{r}<{GL}+{H}$B$19)'
    yr4 = f'AND(A{r}>={GL}+36,A{r}<{GL}+{H}$B$19)'
    put(ws, f'A{r}', m, fmt='0')
    put(ws, f'B{r}', f'=IF({act},{S_}$C$4,0)+IF({yr4},{S_}$C$5,0)+IF(A{r}={GL},{S_}$C$6,0)', fmt=ILS)
    put(ws, f'C{r}', f'=IF({act},{S_}$D$4,0)+IF({yr4},{S_}$D$5,0)', fmt=ILS)
    put(ws, f'D{r}', f'=IF(A{r}=1,{S_}$D$6,0)', fmt=ILS)
    put(ws, f'E{r}', f'=B{r}-C{r}-D{r}', fmt=ILS)
    put(ws, f'F{r}', f'=IF(A{r}-{H}$B$13>=0,INDEX($B${R1}:$B${LAST},A{r}-{H}$B$13+1),0)', fmt=ILS)
    put(ws, f'G{r}', f'=F{r}-C{r}-D{r}', fmt=ILS)
    put(ws, f'H{r}', f'=G{r}' if m == 0 else f'=H{r-1}+G{r}', fmt=ILS)
put(ws, 'A6', 'הנחה: החומרה ושעות ההקמה משולמות לספקים בחודש 1, והחשבון החד-פעמי לעירייה מוגש בחודש ההתקנה. תנאי תשלום בפועל – לפי חוברת המכרז.')
ws.merge_cells('A6:H6')
ws.freeze_panes = f'B{R1}'

# ------------------------------------------------------------------ שגיאות בטופס
ws = sheet('שגיאות בטופס', {'A': 6, 'B': 10, 'C': 70, 'D': 50})
ws['A1'] = 'ליקויים בכתב הכמויות'
ws['A1'].font = TITLE
header(ws, 3, ['#', 'סעיף', 'ליקוי', 'השפעה'])
errs = [
    ('14', 'אין נוסחת סכום (G16 ריק): 24 קווים × 300 ₪ לא נכללים בסה"כ החודשי.', '7,200 ₪ לחודש חסרים בהשוואה.'),
    ('סה"כ', 'שורת "סה"כ לפני מע"מ" ריקה, בלי נוסחה. לא מוגדר אם ההשוואה חודשית, חד-פעמית או כוללת.', 'שיטת ההשוואה לא ידועה.'),
    ('5–8, 29', 'המחיר החד-פעמי נרשם כטקסט ("250 לשעה") בעמודת התשלום החד-פעמי; נוסחת הסה"כ כופלת את העמודה החודשית (ריקה) ומחזירה 0.', 'סעיפי השעות לא נכנסים לסה"כ.'),
    ('23–28', 'כמות כטקסט ("15 יח\'") ומחיר כטקסט ("4700 ליחידה"); אין נוסחאות סכום לחד-פעמי (עמודה H ריקה).', 'כל החומרה (~246,500 ₪) מחוץ לסה"כ.'),
    ('20', 'התקנה חד-פעמית רשומה בעמודה החודשית (1,500 ₪ לחודש).', 'מעוות את הסה"כ החודשי.'),
    ('9', 'תחזוקה "החל משנה רביעית" נסכמת בסה"כ החודשי מהחודש הראשון (ערך קבוע 2,000).', 'מנפח את ההשוואה החודשית.'),
    ('4', 'סה"כ רשום כערך קבוע (3,500) ולא כנוסחה.', 'לא יתעדכן אם הכמות תשתנה.'),
    ('22', 'כמות לא צוינה.', 'לא ידוע אם נכלל בהשוואה.'),
    ('19, 20', 'עמודת היחידה ריקה.', ''),
    ('23', 'יחידה "קו" למתג.', 'טעות סופר.'),
    ('12 מול 10', 'קו 300/150M זול יותר (1,600) מקו 200M (1,850).', 'לבדוק מפרט.'),
]
for i, (s, e, imp) in enumerate(errs):
    r = 4 + i
    put(ws, f'A{r}', i + 1)
    put(ws, f'B{r}', s)
    put(ws, f'C{r}', e).alignment = WRAP
    put(ws, f'D{r}', imp).alignment = WRAP
    ws.row_dimensions[r].height = 32

wb.save(OUT)
print('saved', OUT)
